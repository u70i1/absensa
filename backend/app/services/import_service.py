"""Read one workbook, review changes, then apply selected rows atomically."""

from collections import Counter
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from secrets import token_urlsafe
from typing import BinaryIO
from zipfile import ZipFile

from app.models.class_ import Class
from app.models.import_batch import ImportBatch
from app.models.student import Student
from app.schemas.import_ import ClassImportRow, StudentImportRow
from app.services import class_service, student_service
from app.services.exceptions import AppException
from app.services.export_service import CLASS_COLUMNS, STUDENT_COLUMNS
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import ValidationError
from sqlalchemy import delete, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10000
COLUMNS = {"students": STUDENT_COLUMNS, "classes": CLASS_COLUMNS}
FIELD_ERRORS = {
    "id": "ID harus berupa bilangan bulat positif, atau kosong untuk data baru.",
    "class_id": "ID KELAS harus berupa bilangan bulat positif, atau kosong jika belum ditentukan.",
    "name": "NAMA wajib diisi, maksimal 255 karakter.",
    "nisn": "NISN wajib berisi tepat 10 digit. Gunakan format teks untuk mempertahankan nol di depan.",
    "current": "STATUS wajib diisi dengan Aktif atau Tidak aktif.",
    "guardian_phone": "NOMOR WALI harus berisi angka saja, diawali 08 atau 628, maksimal 32 digit.",
    "grade": "JENJANG wajib berupa bilangan bulat antara 1 dan 20.",
    "class_name": "NAMA KELAS wajib diisi, maksimal 20 karakter.",
}


def _error(kind, row, field, message):
    column = next(i for i, (_, key) in enumerate(COLUMNS[kind], 1) if key == field)
    return {"cell": f"{get_column_letter(column)}{row}", "sheet": kind, "message": message}


def _value(field, value):
    if isinstance(value, str):
        value = value.strip()
    if value in (None, ""):
        return None
    if field == "current":
        if value not in ("Aktif", "Tidak aktif"):
            raise ValueError(FIELD_ERRORS[field])
        return value == "Aktif"
    if field in {"id", "class_id", "grade"}:
        if isinstance(value, bool):
            raise ValueError(FIELD_ERRORS[field])
        if isinstance(value, float) and value.is_integer():
            return int(value)
        if isinstance(value, int) or (isinstance(value, str) and value.isascii() and value.isdigit()):
            return int(value)
        raise ValueError(FIELD_ERRORS[field])
    if not isinstance(value, str):
        # Never guess zeros already lost by Excel's numeric conversion.
        if field in {"nisn", "guardian_phone"} and isinstance(value, (int, float)) and not isinstance(value, bool):
            if int(value) == value:
                return str(int(value))
        raise ValueError(FIELD_ERRORS[field])
    return value


def _read_rows(content: bytes, kind: str):
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > 50 * 1024 * 1024:
                raise AppException("Isi file terlalu besar setelah dibuka. Pecah data menjadi file yang lebih kecil.", 413)
            if any("vbaproject" in item.filename.lower() for item in entries):
                raise AppException("Gunakan file .xlsx tanpa makro.", 422)
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except AppException:
        raise
    except Exception as exc:
        raise AppException("File Excel rusak atau bukan file .xlsx yang valid.", 422) from exc
    try:
        if kind not in workbook.sheetnames:
            raise AppException(f"Sheet {kind} tidak ditemukan. Gunakan template yang sesuai.", 422)
        if any(name in workbook.sheetnames for name in COLUMNS if name != kind):
            raise AppException("Gunakan satu template siswa atau kelas per file.", 422)
        sheet = workbook[kind]
        columns = COLUMNS[kind]
        if (sheet.max_row or 0) > 100000 or (sheet.max_column or 0) > len(columns):
            raise AppException("Ukuran sheet atau jumlah kolom tidak sesuai template.", 422)
        rows = sheet.iter_rows(max_col=len(columns))
        header = next(rows, ())
        if tuple(cell.value for cell in header) != tuple(name for name, _ in columns):
            raise AppException("Judul atau urutan kolom tidak sesuai template. Unduh template terbaru.", 422)
        parsed, errors = [], []
        used_rows = 0
        schema = StudentImportRow if kind == "students" else ClassImportRow
        for number, cells in enumerate(rows, 2):
            if all(cell.value is None or (isinstance(cell.value, str) and not cell.value.strip()) for cell in cells):
                continue
            used_rows += 1
            if used_rows > MAX_ROWS:
                raise AppException("Maksimal 10.000 baris data per file.", 422)
            values, row_errors = {}, []
            for (_, field), cell in zip(columns, cells):
                try:
                    if cell.data_type in {"f", "e"}:
                        raise ValueError("Gunakan nilai langsung, bukan rumus atau nilai kesalahan Excel.")
                    values[field] = _value(field, cell.value)
                except ValueError as exc:
                    row_errors.append(_error(kind, number, field, str(exc)))
            if not row_errors:
                try:
                    values = schema.model_validate(values).model_dump()
                    # StudentBase omits an unspecified optional guardian phone.
                    if kind == "students":
                        values.setdefault("guardian_phone", None)
                except ValidationError as exc:
                    row_errors.extend(_error(kind, number, str(e["loc"][0]), FIELD_ERRORS[str(e["loc"][0])]) for e in exc.errors())
            errors.extend(row_errors)
            if not row_errors:
                parsed.append({"row": number, "values": values})
        if not used_rows:
            raise AppException("File tidak berisi data. Isi baris di bawah judul kolom terlebih dahulu.", 422)
        return parsed, errors, used_rows
    except AppException:
        raise
    except Exception as exc:
        raise AppException("Isi sheet tidak dapat dibaca. Simpan ulang sebagai .xlsx.", 422) from exc
    finally:
        workbook.close()


def _review_rows(db: Session, kind: str, rows: list[dict]):
    id_field = "id" if kind == "students" else "class_id"
    model = Student if kind == "students" else Class
    fields = [field for _, field in COLUMNS[kind] if field != id_field]
    ids = [row["values"][id_field] for row in rows if row["values"][id_field] is not None]
    existing = {getattr(item, id_field): item for item in db.scalars(select(model).where(getattr(model, id_field).in_(ids)))}
    classes = {item.class_id: item for item in db.scalars(select(Class))}
    if kind == "students":
        keys = {item.nisn: item.id for item in db.scalars(select(Student).where(Student.nisn.in_([r["values"]["nisn"] for r in rows])))}
    else:
        keys = {(item.grade, item.class_name): item.class_id for item in classes.values()}
    key_for = lambda values: values["nisn"] if kind == "students" else (values["grade"], values["class_name"])
    key_counts = Counter(key_for(row["values"]) for row in rows)
    id_counts = Counter(ids)
    valid, errors = [], []
    for row in rows:
        values, number = row["values"], row["row"]
        item_id = values[id_field]
        key = key_for(values)
        row_errors = []
        if item_id is not None and item_id not in existing:
            row_errors.append(_error(kind, number, id_field, "ID tidak ditemukan. Untuk membuat data baru, kosongkan ID."))
        if item_id is not None and id_counts[item_id] > 1:
            row_errors.append(_error(kind, number, id_field, "ID muncul lebih dari sekali dalam file."))
        if key_counts[key] > 1 or (key in keys and keys[key] != item_id):
            row_errors.append(_error(kind, number, "nisn" if kind == "students" else "class_name", "NISN sudah digunakan atau berulang dalam file." if kind == "students" else "Kombinasi jenjang dan nama kelas sudah digunakan atau berulang dalam file."))
        if kind == "students" and values["class_id"] is not None and values["class_id"] not in classes:
            row_errors.append(_error(kind, number, "class_id", "Kelas tidak ditemukan. Gunakan ID kelas yang sudah terdaftar."))
        errors.extend(row_errors)
        if row_errors:
            continue
        before = {field: getattr(existing[item_id], field) for field in fields} if item_id else None
        changed = [field for field in fields if before is None or values[field] != before[field]]
        row = {**row, "before": before, "changed": changed, "action": "create" if item_id is None else "update"}
        if kind == "students":
            class_ = classes.get(values["class_id"])
            row["class_name"] = class_.class_name if class_ else "—"
            row["photo_path"] = existing[item_id].photo_path if item_id else None
        valid.append(row)
    return valid, errors


def create_preview(db: Session, admin_id: int, kind: str, filename: str, source: BinaryIO) -> ImportBatch:
    if kind not in COLUMNS:
        raise AppException("Jenis impor tidak ditemukan.", 404)
    if Path(filename).suffix.lower() != ".xlsx":
        raise AppException("Pilih satu file dengan format .xlsx.", 422)
    content = source.read(MAX_BYTES + 1)
    if len(content) > MAX_BYTES:
        raise AppException("Ukuran file maksimal 10 MB.", 413)
    parsed, errors, total = _read_rows(content, kind)
    valid, review_errors = _review_rows(db, kind, parsed)
    now = datetime.now(timezone.utc)
    db.execute(delete(ImportBatch).where(ImportBatch.expires_at < now))
    batch = ImportBatch(
        token=token_urlsafe(32), admin_id=admin_id, kind=kind,
        filename=filename.replace("\\", "/").rsplit("/", 1)[-1][:255], size=len(content),
        payload={"rows": valid, "errors": errors + review_errors, "total": total},
        state="pending", expires_at=now + timedelta(hours=1),
    )
    db.add(batch)
    db.commit()
    return batch


def get_preview(db: Session, admin_id: int, token: str, *, lock=False) -> ImportBatch:
    stmt = select(ImportBatch).where(ImportBatch.token == token, ImportBatch.admin_id == admin_id)
    if lock:
        stmt = stmt.with_for_update().execution_options(populate_existing=True)
    batch = db.scalar(stmt)
    if batch is None or batch.expires_at <= datetime.now(timezone.utc):
        raise AppException("Pratinjau tidak ditemukan atau sudah kedaluwarsa. Unggah file kembali.", 404)
    return batch


def cancel_preview(db: Session, admin_id: int, token: str):
    batch = get_preview(db, admin_id, token, lock=True)
    if batch.state == "pending":
        batch.state = "cancelled"
        batch.payload = {}
        db.commit()


def apply_preview(db: Session, admin_id: int, token: str, selected: list[int]) -> ImportBatch:
    batch = get_preview(db, admin_id, token, lock=True)
    if batch.state == "applied":
        return batch  # A retry never creates duplicate records.
    if batch.state != "pending":
        raise AppException("Impor ini sudah dibatalkan. Unggah file kembali.", 409)
    rows = batch.payload["rows"]
    selected_ids = set(selected)
    if not selected_ids or not selected_ids.issubset({row["row"] for row in rows}):
        raise AppException("Pilih setidaknya satu baris valid dari pratinjau.", 422)
    chosen = [row for row in rows if row["row"] in selected_ids]
    kind = batch.kind
    id_field, model = ("id", Student) if kind == "students" else ("class_id", Class)
    try:
        # Lock existing rows in a stable order and reject stale previews.
        ids = sorted(row["values"][id_field] for row in chosen if row["before"] is not None)
        locked = {getattr(item, id_field): item for item in db.scalars(select(model).where(getattr(model, id_field).in_(ids)).order_by(getattr(model, id_field)).with_for_update().execution_options(populate_existing=True))}
        for row in chosen:
            if row["before"] is not None:
                item = locked.get(row["values"][id_field])
                if item is None or any(getattr(item, field) != value for field, value in row["before"].items()):
                    raise AppException("Data berubah sejak pratinjau dibuat. Batalkan lalu unggah kembali untuk meninjau perubahan terbaru.", 409)
        _, errors = _review_rows(db, kind, chosen)
        if errors:
            raise AppException("Data tidak lagi valid: " + errors[0]["message"] + " Unggah file kembali.", 409)
        for row in chosen:
            values = dict(row["values"])
            item_id = values.pop(id_field)
            if kind == "students":
                if item_id is None:
                    student_service.post_student(db, **values, commit=False)
                else:
                    student_service.edit_student(db, item_id, **values, commit=False)
            elif item_id is None:
                class_service.post_class(db, **values, commit=False)
            else:
                class_service.update_class(db, item_id, **values, commit=False)
        batch.state = "applied"
        batch.payload = {"created": sum(row["action"] == "create" for row in chosen), "updated": sum(row["action"] == "update" for row in chosen), "skipped": batch.payload["total"] - len(chosen)}
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise AppException("Data bertabrakan dengan perubahan lain. Tidak ada perubahan disimpan. Unggah file kembali.", 409) from exc
    except Exception:
        db.rollback()
        raise
    return batch
