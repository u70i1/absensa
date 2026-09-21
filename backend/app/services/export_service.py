"""Build spreadsheet exports from the versioned workbook templates."""

from copy import copy
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

from app.core.config import settings
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

TEMPLATE_DIR = Path(__file__).resolve().parents[1] / "export_templates"
STUDENT_TEMPLATE = TEMPLATE_DIR / "students_export_template.xlsx"
STUDENT_COLUMNS = (
    ("ID", "id"),
    ("NAMA", "name"),
    ("ID KELAS", "class_id"),
    ("NISN", "nisn"),
    ("STATUS", "current"),
    ("NOMOR WALI", "guardian_phone"),
)
CLASS_TEMPLATE = TEMPLATE_DIR / "classes_export_template.xlsx"
CLASS_COLUMNS = (
    ("ID KELAS", "class_id"),
    ("JENJANG", "grade"),
    ("NAMA KELAS", "class_name"),
)


def _copy_cell_style(source, target) -> None:
    target._style = copy(source._style)
    target.number_format = source.number_format
    target.protection = copy(source.protection)
    target.alignment = copy(source.alignment)


def _replace_metadata(workbook, replacements: dict[str, object]) -> None:
    sheet = workbook["Tentang Ekspor"]
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value in replacements:
                cell.value = replacements[cell.value]


def build_students_workbook(students, filter_summary: str) -> bytes:
    """Populate the student template while retaining its styles and validation."""
    workbook = load_workbook(STUDENT_TEMPLATE)
    sheet = workbook["students"]
    students = list(students)
    expected_headers = tuple(header for header, _ in STUDENT_COLUMNS)
    actual_headers = tuple(
        sheet.cell(row=1, column=column).value
        for column in range(1, len(STUDENT_COLUMNS) + 1)
    )
    if actual_headers != expected_headers:
        raise ValueError(
            "Student export template columns do not match the configured mapping: "
            f"expected {expected_headers}, found {actual_headers}."
        )

    status_column = next(
        index
        for index, (_, field) in enumerate(STUDENT_COLUMNS, start=1)
        if field == "current"
    )
    status_cell = f"{get_column_letter(status_column)}2"
    for validation in sheet.data_validations.dataValidation:
        if status_cell in validation.sqref:
            validation.formula1 = '"Aktif,Tidak aktif"'

    exported_at = datetime.now(ZoneInfo(settings.timezone))
    _replace_metadata(
        workbook,
        {
            "{{ filter_summary }}": filter_summary,
            "{{ exported_at }}": exported_at.strftime("%Y-%m-%d %H:%M:%S %Z"),
            "{{ row_count }}": len(students),
        },
    )

    style_cells = tuple(
        sheet.cell(row=2, column=column)
        for column in range(1, len(STUDENT_COLUMNS) + 1)
    )
    for row_number, student in enumerate(students, start=2):
        if row_number > 2:
            sheet.row_dimensions[row_number].height = sheet.row_dimensions[2].height
        for column, (_, field) in enumerate(STUDENT_COLUMNS, start=1):
            cell = sheet.cell(row=row_number, column=column)
            _copy_cell_style(style_cells[column - 1], cell)
            value = getattr(student, field)
            cell.value = (
                ("Aktif" if value else "Tidak aktif") if field == "current" else value
            ) # type: ignore
            if isinstance(cell.value, str):
                cell.data_type = "s"

    if not students:
        for column in range(1, len(STUDENT_COLUMNS) + 1):
            sheet.cell(row=2, column=column).value = None

    last_row = max(2, len(students) + 1)
    last_column = get_column_letter(len(STUDENT_COLUMNS))
    sheet.tables["StudentsExportTable"].ref = f"A1:{last_column}{last_row}"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_classes_workbook(classes, filter_summary: str) -> bytes:
    """Populate the class template while retaining its styles and validation."""
    workbook = load_workbook(CLASS_TEMPLATE)
    sheet = workbook["classes"]
    classes = list(classes)
    expected_headers = tuple(header for header, _ in CLASS_COLUMNS)
    actual_headers = tuple(
        sheet.cell(row=1, column=column).value
        for column in range(1, len(CLASS_COLUMNS) + 1)
    )
    if actual_headers != expected_headers:
        raise ValueError(
            "Class export template columns do not match the configured mapping: "
            f"expected {expected_headers}, found {actual_headers}."
        )

    exported_at = datetime.now(ZoneInfo(settings.timezone))
    _replace_metadata(
        workbook,
        {
            "{{ filter_summary }}": filter_summary,
            "{{ exported_at }}": exported_at.strftime("%Y-%m-%d %H:%M:%S %Z"),
            "{{ row_count }}": len(classes),
        },
    )

    style_cells = tuple(
        sheet.cell(row=2, column=column)
        for column in range(1, len(CLASS_COLUMNS) + 1)
    )
    for row_number, class_ in enumerate(classes, start=2):
        if row_number > 2:
            sheet.row_dimensions[row_number].height = sheet.row_dimensions[2].height
        for column, (_, field) in enumerate(CLASS_COLUMNS, start=1):
            cell = sheet.cell(row=row_number, column=column)
            _copy_cell_style(style_cells[column - 1], cell)
            cell.value = getattr(class_, field)
            if isinstance(cell.value, str):
                cell.data_type = "s"

    if not classes:
        for column in range(1, len(CLASS_COLUMNS) + 1):
            sheet.cell(row=2, column=column).value = None

    last_row = max(2, len(classes) + 1)
    last_column = get_column_letter(len(CLASS_COLUMNS))
    sheet.tables["ClassesExportTable"].ref = f"A1:{last_column}{last_row}"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
