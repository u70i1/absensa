"""Shared class declarations resolve to one class during workbook submission."""

import re
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from app.models.class_ import Class
from app.models.student import Student
from app.services import export_service, import_service
from app.services.exceptions import AppException
from openpyxl import load_workbook
from sqlalchemy import select

from tests.test_import import importer as import_admin
from tests.test_import import selection

importer = import_admin


def workbook(classes, students=()):
    book = load_workbook(
        BytesIO(export_service.build_students_workbook([], "test", []))
    )
    # Remove reserved slots so tests can declare arbitrary workbook-local IDs.
    for sheet in (book["classes"], book["students"]):
        sheet.delete_rows(2, sheet.max_row - 1)
    for values in classes:
        book["classes"].append(values)
    for number, (reference, nisn) in enumerate(students, 2):
        values = {
            "id": None,
            "name": f"Student {nisn}",
            "nisn": nisn,
            "current": "Aktif",
            "guardian_phone": None,
            "class_id": reference,
            "grade": export_service.class_lookup_formula(number, 2),
            "class_name": export_service.class_lookup_formula(number, 3),
        }
        book["students"].append(
            [values[field] for _, field in export_service.STUDENT_COLUMNS]
        )
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def preview(db, admin, files):
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return import_service.create_preview(
        db, admin.id, None, "workbooks.zip", BytesIO(output.getvalue())
    )


def apply(db, admin, batch):
    return import_service.apply_preview(db, admin.id, batch.token, selection(batch))


@pytest.mark.parametrize("references", [("N1", "N1"), ("N1", "N8"), (None, None)])
def test_shared_new_class_submits_once_and_reports_reuse(
    client, db_session, importer, references
):
    first, second = references
    batch = preview(
        db_session,
        importer,
        {
            "one.xlsx": workbook(
                [(first, 10, "10A")], [(first, "0000000001")] if first else []
            ),
            "two.xlsx": workbook(
                [(second, 10, "10A")], [(second, "0000000002")] if second else []
            ),
        },
    )
    assert not batch.payload["errors"]
    assert not list(db_session.scalars(select(Class)))
    class_rows = [r for r in batch.payload["rows"] if r["kind"] == "classes"]
    assert [r["action"] for r in class_rows] == ["create", "reuse"]
    page = client.get(f"/admin/import?batch={batch.token}")
    assert "Gunakan kelas yang sama" in page.text
    assert {
        int(key) for key in re.findall(r'name="selected" value="(\d+)"', page.text)
    } == set(selection(batch))
    keys = selection(batch)
    response = client.post(
        f"/admin/import/{batch.token}/confirm",
        data={"selected": keys},
        follow_redirects=False,
    )
    assert response.status_code == 303, response.text
    classes = list(db_session.scalars(select(Class)))
    students = list(db_session.scalars(select(Student)))
    assert len(classes) == 1
    assert len(students) == (2 if first else 0)
    assert all(s.class_id == classes[0].class_id for s in students)
    assert batch.payload["created"] == len(students) + 1
    assert batch.payload["reused_classes"] == 1
    assert batch.payload["skipped"] == 0
    assert (
        "1 baris kelas otomatis menggunakan kelas yang sama"
        in client.get(response.headers["location"]).text
    )
    # A retry preserves both records and the receipt.
    receipt = dict(batch.payload)
    import_service.apply_preview(db_session, importer.id, batch.token, keys)
    assert batch.payload == receipt
    assert len(list(db_session.scalars(select(Class)))) == 1


@pytest.mark.parametrize("same_workbook", [False, True])
def test_distinct_temporary_references_can_share_one_class(
    db_session, importer, same_workbook
):
    declarations = [("N1", 10, "10A"), ("N2", 10, "10A")]
    students = [("N1", "0000000001"), ("N2", "0000000002")]
    files = (
        {"one.xlsx": workbook(declarations, students)}
        if same_workbook
        else {
            "one.xlsx": workbook(declarations[:1], students[:1]),
            "two.xlsx": workbook(declarations[1:], students[1:]),
        }
    )
    batch = preview(db_session, importer, files)
    assert not batch.payload["errors"]
    apply(db_session, importer, batch)
    classes = list(db_session.scalars(select(Class)))
    assert len(classes) == 1
    assert {s.class_id for s in db_session.scalars(select(Student))} == {
        classes[0].class_id
    }


@pytest.mark.parametrize("already_created", [False, True])
def test_shared_class_reused_between_submissions(db_session, importer, already_created):
    first = preview(
        db_session,
        importer,
        {"one.xlsx": workbook([("N1", 10, "10A")], [("N1", "0000000001")])},
    )
    if already_created:
        apply(db_session, importer, first)
    second = preview(
        db_session,
        importer,
        {"two.xlsx": workbook([("N1", 10, "10A")], [("N1", "0000000002")])},
    )
    assert not second.payload["errors"]
    if not already_created:
        # The other submission creates the class after this preview was made.
        apply(db_session, importer, first)
    apply(db_session, importer, second)
    assert len(list(db_session.scalars(select(Class)))) == 1
    assert len({s.class_id for s in db_session.scalars(select(Student))}) == 1
    assert second.payload["created"] == 1
    assert second.payload["reused_classes"] == 1


@pytest.mark.parametrize("second_pair", [(11, "10A"), (10, "10B")])
def test_temporary_ids_remain_local_and_distinct_pairs_are_not_merged(
    db_session, importer, second_pair
):
    batch = preview(
        db_session,
        importer,
        {
            "one.xlsx": workbook([("N1", 10, "10A")], [("N1", "0000000001")]),
            "two.xlsx": workbook([("N1", *second_pair)], [("N1", "0000000002")]),
        },
    )
    assert not batch.payload["errors"]
    apply(db_session, importer, batch)
    students = list(db_session.scalars(select(Student).order_by(Student.nisn)))
    assert len(list(db_session.scalars(select(Class)))) == 2
    assert students[0].class_id != students[1].class_id
    assert (students[1].class_.grade, students[1].class_.class_name) == second_pair


@pytest.mark.parametrize("rename", [False, True])
def test_matching_existing_class_declarations_across_workbooks(
    db_session, importer, class_factory, rename
):
    existing = class_factory(grade=10, class_name="Old")
    name = "Renamed" if rename else "Old"
    batch = preview(
        db_session,
        importer,
        {
            "one.xlsx": workbook(
                [(existing.class_id, 10, name)], [(existing.class_id, "0000000001")]
            ),
            "two.xlsx": workbook(
                [(existing.class_id, 10, name)], [(existing.class_id, "0000000002")]
            ),
        },
    )
    assert not batch.payload["errors"]
    apply(db_session, importer, batch)
    db_session.refresh(existing)
    assert existing.class_name == name
    assert len(list(db_session.scalars(select(Class)))) == 1
    assert {s.class_id for s in db_session.scalars(select(Student))} == {
        existing.class_id
    }
    assert batch.payload["created"] == 2
    assert batch.payload["updated"] == 1
    assert batch.payload["reused_classes"] == 1


def test_new_declaration_reuses_final_renamed_class_even_when_listed_first(
    db_session, importer, class_factory
):
    existing = class_factory(grade=10, class_name="Old")
    batch = preview(
        db_session,
        importer,
        {
            "one.xlsx": workbook([("N1", 10, "New")], [("N1", "0000000001")]),
            "two.xlsx": workbook([(existing.class_id, 10, "New")]),
        },
    )
    assert not batch.payload["errors"]
    apply(db_session, importer, batch)
    db_session.refresh(existing)
    assert existing.class_name == "New"
    assert len(list(db_session.scalars(select(Class)))) == 1
    assert db_session.scalar(select(Student)).class_id == existing.class_id


@pytest.mark.parametrize("edited_file", ["one.xlsx", "two.xlsx"])
def test_editing_a_shared_class_draft_updates_reuse_and_student_resolution(
    db_session, importer, edited_file
):
    batch = preview(
        db_session,
        importer,
        {
            "one.xlsx": workbook([("N1", 10, "10A")], [("N1", "0000000001")]),
            "two.xlsx": workbook([("N1", 10, "10A")], [("N1", "0000000002")]),
        },
    )
    edited = next(
        r
        for r in batch.payload["rows"]
        if r["kind"] == "classes" and r["file"] == edited_file
    )
    import_service.edit_preview_row(
        db_session,
        importer.id,
        batch.token,
        edited["key"],
        {"grade": "10", "class_name": "10B"},
        0,
    )
    class_rows = [r for r in batch.payload["rows"] if r["kind"] == "classes"]
    assert [r["action"] for r in class_rows] == ["create", "create"]
    apply(db_session, importer, batch)
    expected_nisn = "0000000001" if edited_file == "one.xlsx" else "0000000002"
    students = list(db_session.scalars(select(Student)))
    assert len(students) == 2 and len(list(db_session.scalars(select(Class)))) == 2
    assert (
        next(s for s in students if s.nisn == expected_nisn).class_.class_name == "10B"
    )


@pytest.mark.parametrize(
    "failure", ["conflicting_id", "repeated_id", "duplicate_nisn", "rename_collision"]
)
def test_deduplication_preserves_validation_and_atomic_rejection(
    db_session, importer, class_factory, failure
):
    existing = class_factory(grade=10, class_name="10A")
    if failure == "conflicting_id":
        files = {
            "one.xlsx": workbook([(existing.class_id, 10, "First")]),
            "two.xlsx": workbook([(existing.class_id, 10, "Second")]),
        }
    elif failure == "repeated_id":
        files = {"one.xlsx": workbook([("N1", 10, "New"), ("N1", 10, "New")])}
    elif failure == "duplicate_nisn":
        content = workbook([("N1", 10, "New")], [("N1", "0000000001")])
        files = {"one.xlsx": content, "two.xlsx": content}
    else:
        other = class_factory(grade=10, class_name="10B")
        files = {"one.xlsx": workbook([(other.class_id, 10, "10A")])}
    before = [
        (c.class_id, c.grade, c.class_name) for c in db_session.scalars(select(Class))
    ]
    batch = preview(db_session, importer, files)
    assert batch.payload["errors"]
    with pytest.raises(AppException) as rejected:
        apply(db_session, importer, batch)
    assert rejected.value.status_code == 422
    assert [
        (c.class_id, c.grade, c.class_name) for c in db_session.scalars(select(Class))
    ] == before
    assert not list(db_session.scalars(select(Student)))


def test_shared_class_submission_rolls_back_and_can_retry(
    db_session, importer, monkeypatch
):
    batch = preview(
        db_session,
        importer,
        {
            "one.xlsx": workbook([("N1", 10, "10A")], [("N1", "0000000001")]),
            "two.xlsx": workbook([("N1", 10, "10A")], [("N1", "0000000002")]),
        },
    )
    original = import_service._apply_student

    def fail_second(db, student_id, values):
        if values["nisn"] == "0000000002":
            raise AppException("Test write failure", 503)
        return original(db, student_id, values)

    with monkeypatch.context() as patch:
        patch.setattr(import_service, "_apply_student", fail_second)
        with pytest.raises(AppException):
            apply(db_session, importer, batch)
    assert not list(db_session.scalars(select(Class)))
    assert not list(db_session.scalars(select(Student)))
    db_session.refresh(batch)
    assert batch.state == "pending"
    apply(db_session, importer, batch)
    assert len(list(db_session.scalars(select(Class)))) == 1
    assert len(list(db_session.scalars(select(Student)))) == 2
