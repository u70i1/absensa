"""Administrator attendance history, deletion, and export."""

from datetime import date, datetime, timedelta
from html import unescape
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from app.core.admin_auth import ADMIN_SESSION_COOKIE
from app.core.config import settings
from app.models.admin import Admin
from app.models.scan_log import ScanLog
from app.services.admin_auth_service import create_admin_session, hash_password
from app.services.export_service import STUDENT_TEMPLATE
from openpyxl import load_workbook

SCHOOL_TZ = ZoneInfo(settings.timezone)
DAY = date(2026, 9, 25)


@pytest.fixture
def admin_client(client, db_session):
    admin = Admin(username="scan-admin", password_hash=hash_password("test-password"))
    db_session.add(admin)
    db_session.commit()
    client.cookies.set(
        ADMIN_SESSION_COOKIE, create_admin_session(db_session, admin, 1), path="/"
    )
    return client


def test_log_page_requires_admin(client):
    for path in ("/admin/scans", "/admin/scans/export"):
        response = client.get(path, follow_redirects=False)
        assert response.status_code == 303
        assert response.headers["location"] == "/admin"


def test_log_page_defaults_to_school_today(admin_client):
    response = admin_client.get("/admin/scans")
    assert response.status_code == 200
    assert f'value="{datetime.now(SCHOOL_TZ).date().isoformat()}"' in response.text


def test_day_and_name_filter_use_saved_scan_name(
    admin_client, db_session, student_factory, scan_log_factory
):
    student = student_factory(name="Nadia")
    morning = scan_log_factory(student, datetime(2026, 9, 25, 0, 10, tzinfo=SCHOOL_TZ))
    other_day = scan_log_factory(
        student, datetime(2026, 9, 26, 0, 10, tzinfo=SCHOOL_TZ)
    )
    db_session.delete(student)
    db_session.commit()

    response = admin_client.get(
        "/admin/scans", params={"day": DAY.isoformat(), "q": "nad"}
    )
    assert response.status_code == 200
    assert f'<td class="scan-id">{morning.scan_id}</td>' in response.text
    assert f'<td class="scan-id">{other_day.scan_id}</td>' not in response.text
    assert "Nadia" in response.text
    assert "Log Presensi" in response.text
    assert 'name="day"' in response.text
    assert 'data-selection-toggle' in response.text
    assert "/admin/scans/export?" in response.text
    navigation = unescape(response.text)
    assert "/admin/scans?day=2026-09-24&q=nad&page=1" in navigation
    assert "/admin/scans?day=2026-09-26&q=nad&page=1" in navigation
    assert "Tanggal sebelumnya, 24/09/2026" in response.text
    assert "Tanggal berikutnya, 26/09/2026" in response.text
    assert (
        admin_client.get("/admin/scans", params={"day": "yesterday"}).status_code
        == 422
    )


def test_htmx_filter_and_pagination(admin_client, db_session, student_factory):
    student = student_factory()
    for index in range(51):
        db_session.add(
            ScanLog(
                student_id=student.id,
                name=f"Siswa {index:02}",
                timestamp=datetime(2026, 9, 25, 8, tzinfo=SCHOOL_TZ)
                + timedelta(seconds=index),
            )
        )
    db_session.commit()
    first_page = admin_client.get(
        "/admin/scans", params={"day": DAY.isoformat(), "limit": 1}
    )
    assert first_page.status_code == 200
    assert first_page.text.count('class="student-row scan-row"') == 50
    assert "Siswa 50" in first_page.text
    assert "Siswa 00" not in first_page.text
    response = admin_client.get(
        "/admin/scans",
        params={"day": DAY.isoformat(), "page": 2},
        headers={"HX-Request": "true"},
    )
    assert response.status_code == 200
    assert response.text.count('class="student-row scan-row"') == 1
    assert 'id="scan-results"' in response.text
    assert "Siswa 00" in response.text
    assert "Siswa 50" not in response.text
    assert response.headers["HX-Push-Url"].endswith("page=2")
    export = admin_client.get("/admin/scans/export", params={"day": DAY.isoformat()})
    workbook = load_workbook(BytesIO(export.content))
    assert workbook["scan_logs"].tables["ScansExportTable"].ref == "A1:F52"
    assert workbook["Tentang Ekspor"]["D8"].value == 51


def test_row_and_bulk_delete_require_confirmation(
    admin_client, db_session, student_factory, scan_log_factory
):
    student = student_factory()
    first = scan_log_factory(student, datetime(2026, 9, 25, 8, tzinfo=SCHOOL_TZ))
    second = scan_log_factory(student, datetime(2026, 9, 25, 9, tzinfo=SCHOOL_TZ))
    ids = [first.scan_id, second.scan_id]
    page = admin_client.get("/admin/scans", params={"day": DAY.isoformat()})
    assert f'value="{first.scan_id}"' in page.text
    confirmation = admin_client.post(
        "/admin/scans/selection/confirm",
        params={"day": DAY.isoformat()},
        data={"action": "delete", "ids": [str(i) for i in ids]},
        headers={"HX-Request": "true"},
    )
    assert confirmation.status_code == 200
    assert "Hapus 2 log presensi?" in confirmation.text
    assert all(db_session.get(ScanLog, scan_id) for scan_id in ids)

    applied = admin_client.post(
        "/admin/scans/selection/apply",
        params={"day": DAY.isoformat()},
        data={"action": "delete", "ids": [str(first.scan_id)]},
        headers={"HX-Request": "true"},
    )
    assert applied.status_code == 200
    assert applied.headers["HX-Retarget"] == "#scan-results"
    assert db_session.get(ScanLog, first.scan_id) is None
    assert db_session.get(ScanLog, second.scan_id) is not None


def test_scan_selection_rejects_invalid_or_stale_ids(
    admin_client, db_session, student_factory, scan_log_factory
):
    student = student_factory()
    scan = scan_log_factory(student)
    url = "/admin/scans/selection/apply"
    assert (
        admin_client.post(
            url, data={"action": "deactivate", "ids": str(scan.scan_id)}
        ).status_code
        == 422
    )
    assert (
        admin_client.post(
            url, data={"action": "delete", "ids": [str(scan.scan_id), "999999"]}
        ).status_code
        == 409
    )
    assert (
        admin_client.post(
            url,
            data={"action": "delete", "ids": str(scan.scan_id)},
            headers={"Origin": "https://evil.example"},
        ).status_code
        == 403
    )
    assert db_session.get(ScanLog, scan.scan_id) is not None


def test_export_all_filtered_rows_and_not_importable(
    admin_client, student_factory, scan_log_factory
):
    student = student_factory(name="Nadia")
    scan = scan_log_factory(student, datetime(2026, 9, 25, 9, 12, 13, tzinfo=SCHOOL_TZ))
    scan_log_factory(student, datetime(2026, 9, 26, 9, 12, 13, tzinfo=SCHOOL_TZ))
    response = admin_client.get(
        "/admin/scans/export", params={"day": DAY.isoformat(), "q": "nad"}
    )
    assert response.status_code == 200
    assert (
        "ekspor-log-presensi-2026-09-25.xlsx"
        in response.headers["content-disposition"]
    )
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["Tentang Ekspor", "scan_logs"]
    sheet = workbook["scan_logs"]
    assert [sheet.cell(2, col).value for col in range(1, 7)] == [
        scan.scan_id,
        "Nadia",
        None,
        student.nisn,
        datetime(2026, 9, 25),
        datetime(2026, 9, 25, 9, 12, 13).time(),
    ]
    assert sheet.tables["ScansExportTable"].ref == "A1:F2"
    assert not sheet.data_validations.dataValidation
    source = load_workbook(STUDENT_TEMPLATE)
    assert sheet["A1"].fill.fgColor.rgb == source["students"]["A1"].fill.fgColor.rgb
    assert (
        workbook["Tentang Ekspor"]["B2"].fill.fgColor.rgb
        == source["Tentang Ekspor"]["B2"].fill.fgColor.rgb
    )
    assert (
        "tidak dapat digunakan untuk impor" in workbook["Tentang Ekspor"]["B18"].value
    )
    assert workbook["Tentang Ekspor"]["D8"].value == 1


def test_empty_export_has_headers_and_no_import_rows(admin_client):
    response = admin_client.get("/admin/scans/export", params={"day": DAY.isoformat()})
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook["scan_logs"]
    assert sheet.max_row == 2
    assert sheet["A2"].value is None
    assert workbook["Tentang Ekspor"]["D8"].value == 0
    imported = admin_client.post(
        "/admin/import/preview",
        files={"file": ("log-presensi.xlsx", response.content)},
    )
    assert imported.status_code == 422
