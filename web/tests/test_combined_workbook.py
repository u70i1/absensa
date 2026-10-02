"""Stable temporary class references, live lookups and atomic workbook imports."""

from io import BytesIO

import pytest
from app.models.class_ import Class
from app.models.student import Student
from app.services import export_service, import_service
from app.services.exceptions import AppException
from openpyxl import load_workbook
from sqlalchemy import event, select, text

from tests.test_import import importer  # noqa: F401


def save(book):
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def workbook(db, students=()):
    return load_workbook(
        BytesIO(
            export_service.build_students_workbook(
                students,
                "Jenjang 11",
                db.scalars(select(Class).order_by(Class.class_id)),
            )
        )
    )


def student_row(
    sheet, row, *, id=None, name="New student", class_id="N1", nisn="0000000001"
):
    for column, value in enumerate(
        (
            id,
            name,
            nisn,
            "Aktif",
            None,
            export_service.class_lookup_formula(row, 2),
            export_service.class_lookup_formula(row, 3),
            class_id,
        ),
        1,
    ):
        sheet.cell(row, column).value = value


def preview(db, admin, book):
    return import_service.create_preview(
        db, admin.id, None, "combined.xlsx", BytesIO(save(book))
    )


def apply(db, admin, batch):
    return import_service.apply_preview(
        db, admin.id, batch.token, [r["key"] for r in batch.payload["rows"]]
    )


def snapshot(db):
    db.expire_all()
    return (
        list(
            db.execute(
                select(
                    Student.id,
                    Student.name,
                    Student.class_id,
                    Student.nisn,
                    Student.current,
                    Student.guardian_phone,
                ).order_by(Student.id)
            )
        ),
        list(
            db.execute(
                select(Class.class_id, Class.grade, Class.class_name).order_by(
                    Class.class_id
                )
            )
        ),
    )


@pytest.mark.parametrize(
    "mode", ["roundtrip", "rename", "rename_grade", "reassign", "temporary", "unassign"]
)
def test_existing_student_resolution(
    db_session, importer, existing_student, class_factory, mode
):
    db, student = db_session, existing_student
    original_id = student.class_id
    other = class_factory(grade=12, class_name="Other")
    book = workbook(db, [student])
    expected_id = original_id
    if mode.startswith("rename"):
        book["classes"]["C2"] = "Renamed"
        if mode == "rename_grade":
            book["classes"]["B2"] = 9
    elif mode == "reassign":
        book["students"]["H2"] = other.class_id
        expected_id = other.class_id
    elif mode == "temporary":
        assert book["classes"]["A4"].value == "N1"
        book["classes"]["B4"], book["classes"]["C4"] = 6, "New"
        book["students"]["H2"] = "N1"
        expected_id = None
    elif mode == "unassign":
        book["students"]["H2"] = None
        expected_id = None
    batch = preview(db, importer, book)
    assert not batch.payload["errors"]
    apply(db, importer, batch)
    db.refresh(student)
    if mode == "temporary":
        assert isinstance(student.class_id, int)
        assigned = db.get(Class, student.class_id)
        assert (assigned.grade, assigned.class_name) == (6, "New")
    else:
        assert student.class_id == expected_id
    if mode.startswith("rename"):
        assert db.get(Class, original_id).class_name == "Renamed"
    fresh = workbook(db, [student])
    assert fresh["students"]["H2"].value == student.class_id
    assert fresh["students"]["G2"].data_type == "f"
    assert len(list(db.scalars(select(Class)))) == (3 if mode == "temporary" else 2)


@pytest.mark.parametrize("count", [1, 3])
def test_temporary_class_created_once(db_session, importer, count):
    book = workbook(db_session)
    book["classes"]["B2"], book["classes"]["C2"] = 10, "X RPL 1"
    for row in range(2, count + 2):
        student_row(book["students"], row, nisn=f"{row:010}")
    batch = preview(db_session, importer, book)
    assert not batch.payload["errors"]
    apply(db_session, importer, batch)
    classes = list(db_session.scalars(select(Class)))
    students = list(db_session.scalars(select(Student)))
    assert len(classes) == 1 and len(students) == count
    assert {s.class_id for s in students} == {classes[0].class_id}


@pytest.mark.parametrize(
    "failure,cell",
    [
        ("unknown_numeric", "2:H"),
        ("unknown_temporary", "2:H"),
        ("unused_slot", "2:H"),
        ("incomplete_class", "3:C"),
        ("bad_temp_format", "2:H"),
        ("missing_new_reference", "3:H"),
        ("duplicate_class_id", "3:A"),
        ("duplicate_temp_id", "4:A"),
        ("unknown_class_id", "2:A"),
        ("invalid_class", "2:B"),
        ("invalid_student", "3:C"),
        ("unknown_student", "2:A"),
        ("duplicate_student", "3:A"),
        ("missing_sheet", "1:A"),
        ("renamed_header", "1:G"),
        ("missing_header", "1:H"),
        ("manual_grade", "2:F"),
        ("manual_name", "2:G"),
        ("foreign_formula", "2:F"),
    ],
)
def test_invalid_workbook_has_zero_effect(
    db_session, importer, existing_student, failure, cell
):
    db = db_session
    before = snapshot(db)
    book = workbook(db, [existing_student])
    s, c = book["students"], book["classes"]
    s["B2"] = "Valid student update"
    c["C2"] = "Valid rename"
    c["B3"], c["C3"] = 5, "Valid new class"
    if failure == "unknown_numeric":
        s["H2"] = 2147483647
    elif failure == "unknown_temporary":
        s["H2"] = "N999"
    elif failure == "unused_slot":
        s["H2"] = "N2"
    elif failure == "incomplete_class":
        c["C3"] = None
    elif failure == "bad_temp_format":
        s["H2"] = "N01"
    elif failure == "missing_new_reference":
        student_row(s, 3, class_id=None)
    elif failure == "duplicate_class_id":
        c["A3"] = c["A2"].value
    elif failure == "duplicate_temp_id":
        c["A4"], c["B4"], c["C4"] = "N1", 6, "Another"
    elif failure == "unknown_class_id":
        c["A2"] = 2147483647
    elif failure == "invalid_class":
        c["B2"] = "invalid"
    elif failure == "invalid_student":
        student_row(s, 3, nisn="invalid")
    elif failure == "unknown_student":
        s["A2"] = 2147483647
    elif failure == "duplicate_student":
        student_row(
            s,
            3,
            id=existing_student.id,
            class_id=existing_student.class_id,
            nisn=existing_student.nisn,
        )
    elif failure == "missing_sheet":
        del book["classes"]
    elif failure == "renamed_header":
        s["G1"] = "Other"
    elif failure == "missing_header":
        s["H1"] = None
    elif failure == "manual_grade":
        s["F2"] = 11
    elif failure == "manual_name":
        s["G2"] = "11B"
    elif failure == "foreign_formula":
        s["F2"] = "=11"
    batch = preview(db, importer, book)
    assert cell in {e["cell"] for e in batch.payload["errors"]}
    assert snapshot(db) == before
    with pytest.raises(AppException):
        apply(db, importer, batch)
    assert snapshot(db) == before


def test_temporary_ids_survive_row_operations(db_session, importer):
    book = workbook(db_session)
    c = book["classes"]
    c["B2"], c["C2"] = 10, "First"
    c["B3"], c["C3"] = 11, "Second"
    # Move complete records as a spreadsheet sort would, then insert/delete elsewhere.
    first, second = (
        [c.cell(2, i).value for i in (1, 2, 3)],
        [c.cell(3, i).value for i in (1, 2, 3)],
    )
    for col in (1, 2, 3):
        c.cell(2, col).value, c.cell(3, col).value = second[col - 1], first[col - 1]
    c.insert_rows(2)
    c.delete_rows(5)
    c["C3"] = "Renamed second"
    student_row(book["students"], 2, class_id="N2")
    assert c["A3"].value == "N2" and c["A4"].value == "N1"
    assert c["A3"].data_type == "s"
    batch = preview(db_session, importer, book)
    assert not batch.payload["errors"]
    apply(db_session, importer, batch)
    student = db_session.scalar(select(Student))
    assert db_session.get(Class, student.class_id).class_name == "Renamed second"


@pytest.mark.parametrize("stage", ["student_write", "commit", "constraint"])
def test_runtime_failure_rolls_back_everything(
    db_session, importer, existing_student, monkeypatch, stage
):
    db = db_session
    before = snapshot(db)
    book = workbook(db, [existing_student])
    book["classes"]["C2"] = "Renamed"
    book["classes"]["B3"], book["classes"]["C3"] = 5, "New"
    book["students"]["B2"] = "Changed"
    student_row(book["students"], 3)
    batch = preview(db, importer, book)
    assert not batch.payload["errors"]
    original = import_service._apply_student

    def fail_write(db, student_id, values):
        if stage == "constraint" and student_id is None:
            values["nisn"] = existing_student.nisn
        result = original(db, student_id, values)
        if stage == "student_write" and student_id is None:
            raise AppException("Kegagalan penyimpanan", 503)
        return result

    def fail_commit(session):
        session.execute(text("SELECT 1 / 0"))

    monkeypatch.setattr(import_service, "_apply_student", fail_write)
    if stage == "commit":
        event.listen(db, "before_commit", fail_commit)
    try:
        with pytest.raises(AppException):
            apply(db, importer, batch)
    finally:
        if stage == "commit":
            event.remove(db, "before_commit", fail_commit)
    assert snapshot(db) == before


def test_workbook_formulas_validation_and_protection(
    client, importer, existing_student, class_factory
):
    other = class_factory(grade=3, class_name="Different grade")
    response = client.get("/admin/students/export?grade=11")
    book = load_workbook(BytesIO(response.content))
    s, c = book["students"], book["classes"]
    assert book.sheetnames == ["instructions", "students", "classes"]
    assert s["A2"].value == existing_student.id and s["A3"].value is None
    assert s["H2"].value == existing_student.class_id
    assert not s.column_dimensions["H"].hidden and not s["H2"].protection.locked
    assert s.protection.sheet and c.protection.sheet and c["A2"].protection.locked
    for row in (2, 3, 100):
        for col, lookupcol in ((6, 2), (7, 3)):
            assert s.cell(row, col).value == export_service.class_lookup_formula(
                row, lookupcol
            )
            assert s.cell(row, col).protection.locked
            assert s.cell(row, col).font.italic
    assert any(
        v.type == "custom"
        and "H2" in v.sqref
        and "COUNTIF(class_ids,H2)" in v.formula1
        and v.showErrorMessage
        for v in s.data_validations.dataValidation
    )
    assert not any(
        v.type == "list" and any(f"{col}2" in v.sqref for col in ("F", "G", "H"))
        for v in s.data_validations.dataValidation
    )
    assert book.defined_names["class_data"].attr_text == "'classes'!$A:$C"
    assert book.calculation.fullCalcOnLoad and book.calculation.calcMode == "auto"
    assert {existing_student.class_id, other.class_id} <= {c["A2"].value, c["A3"].value}
    assert c["A4"].value == "N1" and c["A103"].value == "N100"
    assert c["A4"].protection.locked and not c["B4"].protection.locked
    assert s["F2"].fill.fgColor.rgb != s["H2"].fill.fgColor.rgb


def test_deleted_class_row_never_deletes_database_class(
    db_session, importer, existing_student, class_factory
):
    other = class_factory(grade=12, class_name="Keep")
    book = workbook(db_session, [existing_student])
    book["classes"].delete_rows(3)
    batch = preview(db_session, importer, book)
    assert not batch.payload["errors"]
    apply(db_session, importer, batch)
    assert db_session.get(Class, other.class_id).class_name == "Keep"


@pytest.mark.parametrize("spelling", ["FALSE()", "FALSE"])
def test_calc_formula_roundtrip_without_cached_values(
    db_session, importer, existing_student, spelling
):
    book = workbook(db_session, [existing_student])
    for row in book["students"].iter_rows(min_row=2):
        for cell in row[5:7]:
            if cell.data_type == "f":
                cell.value = cell.value.replace("FALSE", spelling)
    batch = preview(db_session, importer, book)
    assert not batch.payload["errors"]
    assert batch.payload["total"] == 2  # Empty formula rows and reserved slots ignored.
    apply(db_session, importer, batch)
    assert db_session.scalar(select(Student)).class_id == existing_student.class_id


def test_preview_editor_assigns_temporary_id(
    db_session, client, importer, existing_student
):
    book = workbook(db_session, [existing_student])
    book["classes"]["B3"], book["classes"]["C3"] = 12, "New"
    batch = preview(db_session, importer, book)
    response = client.post(
        f"/admin/import/{batch.token}/rows/2/edit",
        data={
            "name": existing_student.name,
            "nisn": existing_student.nisn,
            "class_id": "N1",
            "current": "Aktif",
            "revision": 0,
            "grade": 1,
            "class_name": "Ignored independent fields",
        },
    )
    assert response.status_code == 200
    db_session.refresh(batch)
    row = next(row for row in batch.payload["rows"] if row["kind"] == "students")
    assert row["class_name"] == "New" and row["values"]["class_id"] == "N1"
    apply(db_session, importer, batch)
    db_session.refresh(existing_student)
    assert db_session.get(Class, existing_student.class_id).class_name == "New"
