"""Import previews and writes include only actual class/student changes."""

import re
from io import BytesIO

import pytest
from app.core.config import settings
from app.models.import_batch import ImportPhoto
from app.models.student import Student
from app.services import export_service, import_service, student_photo_service
from app.services.exceptions import AppException
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import event, select

from tests.test_import import importer as import_admin
from tests.test_import_class_deduplication import apply, preview, workbook

importer = import_admin


def exported(db, student, *, student_changes=None, class_changes=None):
    book = load_workbook(
        BytesIO(
            export_service.build_students_workbook([student], "test", [student.class_])
        )
    )
    for sheet, columns, changes in (
        ("students", export_service.STUDENT_COLUMNS, student_changes),
        ("classes", export_service.CLASS_COLUMNS, class_changes),
    ):
        for field, value in (changes or {}).items():
            column = next(i for i, (_, key) in enumerate(columns, 1) if key == field)
            book[sheet].cell(2, column).value = value
    # Moving a reserved ID into the first class row must remove its unused slot.
    reference = (class_changes or {}).get("class_id")
    if reference is not None:
        sheet = book["classes"]
        for number in range(sheet.max_row, 2, -1):
            if sheet.cell(number, 1).value == reference and all(
                sheet.cell(number, column).value is None for column in (2, 3)
            ):
                sheet.delete_rows(number)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def reviewed(batch, kind):
    return next(r for r in batch.payload["rows"] if r["kind"] == kind)


def test_unchanged_roundtrip_has_no_edit_rows_or_record_writes(
    client, db_session, importer, existing_student, monkeypatch
):
    batch = preview(
        db_session, importer, {"same.xlsx": exported(db_session, existing_student)}
    )
    assert not batch.payload["errors"]
    assert all(
        r["action"] == "unchanged" and not r["changed"] for r in batch.payload["rows"]
    )
    page = client.get(f"/admin/import?batch={batch.token}")
    assert "Tidak ada perubahan data" in page.text
    assert 'data-action="update"' not in page.text
    assert "data-confirm-import" not in page.text

    def unexpected_write(*args):
        pytest.fail("An unchanged student must not reach the write operation")

    monkeypatch.setattr(import_service, "_apply_student", unexpected_write)
    writes = []

    def track_writes(connection, cursor, statement, parameters, context, executemany):
        if re.match(
            r"\s*(INSERT INTO|UPDATE|DELETE FROM) (students|classes)\b",
            statement,
            re.IGNORECASE,
        ):
            writes.append(statement)

    connection = db_session.connection()
    event.listen(connection, "before_cursor_execute", track_writes)
    try:
        apply(db_session, importer, batch)
    finally:
        event.remove(connection, "before_cursor_execute", track_writes)
    assert writes == []
    assert batch.payload["created"] == batch.payload["updated"] == 0
    assert batch.payload["skipped"] == 2


@pytest.mark.parametrize(
    "field,value",
    [
        ("name", "Changed student"),
        ("nisn", "0000000017"),
        ("current", "Tidak aktif"),
        ("guardian_phone", "081234567890"),
        ("class_id", None),
    ],
)
def test_each_student_field_marks_only_the_changed_student(
    db_session, importer, existing_student, field, value
):
    batch = preview(
        db_session,
        importer,
        {
            "changed.xlsx": exported(
                db_session, existing_student, student_changes={field: value}
            )
        },
    )
    assert not batch.payload["errors"]
    assert reviewed(batch, "classes")["action"] == "unchanged"
    student_row = reviewed(batch, "students")
    assert student_row["action"] == "update" and student_row["changed"] == [field]
    apply(db_session, importer, batch)
    db_session.refresh(existing_student)
    assert getattr(existing_student, field) == (False if field == "current" else value)
    assert batch.payload["updated"] == 1
    assert batch.payload["skipped"] == 1


@pytest.mark.parametrize("field,value", [("class_name", "Renamed"), ("grade", 12)])
def test_class_description_changes_do_not_edit_its_students(
    client, db_session, importer, existing_student, field, value
):
    original_id = existing_student.class_id
    batch = preview(
        db_session,
        importer,
        {
            "changed.xlsx": exported(
                db_session, existing_student, class_changes={field: value}
            )
        },
    )
    assert reviewed(batch, "classes")["changed"] == [field]
    assert reviewed(batch, "students")["action"] == "unchanged"
    assert not reviewed(batch, "students")["changed"]
    page = client.get(f"/admin/import?batch={batch.token}")
    assert len(re.findall(r'data-action="update"', page.text)) == 1
    # Hidden unchanged rows still participate in complete, atomic validation.
    keys = [int(k) for k in re.findall(r'name="selected" value="(\d+)"', page.text)]
    response = client.post(
        f"/admin/import/{batch.token}/confirm",
        data={"selected": keys},
        follow_redirects=False,
    )
    assert response.status_code == 303, response.text
    db_session.refresh(existing_student)
    assert existing_student.class_id == original_id
    assert getattr(existing_student.class_, field) == value
    assert batch.payload["updated"] == 1


@pytest.mark.parametrize("temporary", [False, True])
def test_equivalent_class_references_are_not_student_edits(
    db_session, importer, existing_student, temporary
):
    reference = "N1" if temporary else str(existing_student.class_id)
    batch = preview(
        db_session,
        importer,
        {
            "same.xlsx": exported(
                db_session,
                existing_student,
                student_changes={
                    "class_id": reference,
                    "name": f" {existing_student.name} ",
                    "guardian_phone": "",
                },
                class_changes={"class_id": reference},
            )
        },
    )
    assert not batch.payload["errors"]
    student_row = reviewed(batch, "students")
    assert student_row["resolved_class_id"] == existing_student.class_id
    assert student_row["action"] == "unchanged" and student_row["changed"] == []
    apply(db_session, importer, batch)
    assert batch.payload["updated"] == 0
    assert batch.payload["skipped"] == 2


@pytest.mark.parametrize("kind", ["classes", "students"])
def test_reverting_a_draft_removes_it_from_edits(
    client, db_session, importer, existing_student, kind
):
    batch = preview(
        db_session,
        importer,
        {
            "changed.xlsx": exported(
                db_session,
                existing_student,
                student_changes={"name": "Changed"} if kind == "students" else None,
                class_changes={"class_name": "Changed"} if kind == "classes" else None,
            )
        },
    )
    row = reviewed(batch, kind)
    values = dict(row["values"])
    if kind == "students":
        values.update(name=existing_student.name, current="Aktif")
    else:
        values["class_name"] = existing_student.class_.class_name
    import_service.edit_preview_row(
        db_session, importer.id, batch.token, row["key"], values, 0
    )
    assert reviewed(batch, kind)["action"] == "unchanged"
    assert reviewed(batch, kind)["changed"] == []
    assert (
        'data-action="update"'
        not in client.get(f"/admin/import?batch={batch.token}").text
    )
    apply(db_session, importer, batch)
    assert batch.payload["updated"] == 0


def test_reverting_a_draft_keeps_original_stale_data_checks(
    db_session, importer, existing_student
):
    original_name = existing_student.name
    batch = preview(
        db_session,
        importer,
        {
            "changed.xlsx": exported(
                db_session, existing_student, student_changes={"name": "Draft"}
            )
        },
    )
    row = reviewed(batch, "students")
    existing_student.name = "Saved elsewhere"
    db_session.commit()
    values = {**row["values"], "name": original_name, "current": "Aktif"}
    import_service.edit_preview_row(
        db_session, importer.id, batch.token, row["key"], values, 0
    )
    assert reviewed(batch, "students")["before"]["name"] == original_name
    with pytest.raises(AppException) as rejected:
        apply(db_session, importer, batch)
    assert rejected.value.status_code == 409
    db_session.refresh(existing_student)
    assert existing_student.name == "Saved elsewhere"


@pytest.mark.parametrize("identical", [False, True])
def test_photo_only_changes_are_detected_and_identical_uploads_keep_the_file(
    db_session, importer, existing_student, tmp_path, monkeypatch, identical
):
    monkeypatch.setattr(settings, "photos_dir", str(tmp_path / "photos"))

    def photo(color):
        output = BytesIO()
        Image.new("RGB", (24, 24), color).save(output, "PNG")
        return output.getvalue()

    normalized = student_photo_service.normalize_photo(BytesIO(photo("blue")))
    original_path = student_photo_service.store_normalized_photo(normalized)
    existing_student.photo_path = original_path
    db_session.commit()
    batch = preview(
        db_session,
        importer,
        {
            "data.xlsx": exported(db_session, existing_student),
            f"photos/{existing_student.nisn}.png": photo(
                "blue" if identical else "red"
            ),
        },
    )
    assert not batch.payload["errors"]
    student_row = reviewed(batch, "students")
    assert student_row["action"] == ("unchanged" if identical else "update")
    assert student_row["changed"] == ([] if identical else ["photo_path"])
    assert batch.payload["photo_count"] == int(not identical)
    apply(db_session, importer, batch)
    db_session.refresh(existing_student)
    assert (existing_student.photo_path == original_path) == identical
    assert student_photo_service.photo_file(existing_student.photo_path).is_file()
    assert student_photo_service.photo_file(original_path).exists() == identical
    assert batch.payload["updated"] == int(not identical)
    assert batch.payload["photos_saved"] == int(not identical)
    assert not list(db_session.scalars(select(ImportPhoto)))


def test_mixed_workbooks_only_write_actual_creates_and_class_changes(
    db_session, importer, existing_student
):
    batch = preview(
        db_session,
        importer,
        {
            "existing.xlsx": exported(
                db_session, existing_student, class_changes={"class_name": "Renamed"}
            ),
            "new.xlsx": workbook([("N1", 11, "Renamed")], [("N1", "0000000001")]),
        },
    )
    assert not batch.payload["errors"]
    apply(db_session, importer, batch)
    assert batch.payload["created"] == 1
    assert batch.payload["updated"] == 1
    assert batch.payload["skipped"] == 2
    assert {s.class_id for s in db_session.scalars(select(Student))} == {
        existing_student.class_id
    }
