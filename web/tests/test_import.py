"""Spreadsheet import previews, validation, authorization, and atomic writes."""

from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from app.core.admin_auth import ADMIN_SESSION_COOKIE
from app.models.admin import Admin
from app.models.class_ import Class
from app.models.import_batch import ImportBatch
from app.models.student import Student
from app.services import export_service, import_service, student_service
from app.services.admin_auth_service import create_admin_session
from app.services.exceptions import AppException
from openpyxl import load_workbook
from sqlalchemy import func, select


@pytest.fixture
def importer(client, db_session):
    admin = Admin(username="import-admin", password_hash="unused")
    db_session.add(admin)
    db_session.commit()
    token = create_admin_session(db_session, admin, 1)
    client.cookies.set(ADMIN_SESSION_COOKIE, token, path="/admin")
    return admin


def workbook_bytes(kind, rows):
    builder = (
        export_service.build_students_workbook
        if kind == "students"
        else export_service.build_classes_workbook
    )
    workbook = load_workbook(BytesIO(builder([], "test")))
    sheet = workbook[kind]
    for number, row in enumerate(rows, 2):
        for column, value in enumerate(row, 1):
            sheet.cell(number, column, value)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def preview(client, kind="students", rows=None, content=None, filename="data.xlsx"):
    return client.post(
        f"/admin/import/preview/{kind}",
        files={
            "file": (
                filename,
                content if content is not None else workbook_bytes(kind, rows or []),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        follow_redirects=False,
    )


def batch_for(response, client, db_session):
    assert response.status_code == 303, response.text
    page = client.get(response.headers["location"])
    assert page.status_code == 200, page.text
    token = response.headers["location"].split("batch=")[1].split("#")[0]
    return db_session.get(ImportBatch, token), page


def test_import_requires_admin(client):
    for path in (
        "/admin/import",
        "/admin/import/template/students",
        "/admin/import/template/classes",
    ):
        assert client.get(path, follow_redirects=False).status_code == 303
    assert preview(client).status_code == 303
    assert (
        client.post("/admin/import/missing/confirm", follow_redirects=False).status_code
        == 303
    )


def test_import_page_and_template_downloads(client, importer):
    page = client.get("/admin/import")
    assert page.status_code == 200
    assert "Impor data" in page.text
    assert "multiple" not in page.text
    for kind in ("students", "classes"):
        response = client.get(f"/admin/import/template/{kind}")
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.content))
        assert kind in workbook.sheetnames
        assert workbook[kind]["A2"].value is None


def test_preview_then_selected_create_update_and_retry(
    client, importer, db_session, existing_student
):
    student = existing_student
    response = preview(
        client,
        rows=[
            [
                student.id,
                "Nama Baru",
                student.class_id,
                student.nisn,
                "Tidak aktif",
                "08123456789",
            ],
            [None, "Siswa Baru", None, "0012345678", "Aktif", None],
            [None, "Batalkan Saya", None, "0012345679", "Aktif", None],
        ],
    )
    batch, page = batch_for(response, client, db_session)
    assert student.name != "Nama Baru"
    assert db_session.scalar(select(func.count()).select_from(Student)) == 1
    assert "Konfirmasi perubahan data Siswa" in page.text
    assert "import-changed" in page.text
    response = client.post(
        f"/admin/import/{batch.token}/confirm",
        data={"selected": [2, 3]},
        follow_redirects=False,
    )
    assert response.status_code == 303
    db_session.refresh(student)
    assert student.name == "Nama Baru" and student.current is False
    assert student.guardian_phone == "08123456789"
    assert (
        db_session.scalar(select(Student).where(Student.nisn == "0012345678"))
        is not None
    )
    assert db_session.scalar(select(func.count()).select_from(Student)) == 2
    assert "Impor berhasil" in client.get(response.headers["location"]).text
    assert (
        client.post(
            f"/admin/import/{batch.token}/confirm", data={"selected": [2, 3]}
        ).status_code
        == 200
    )
    assert db_session.scalar(select(func.count()).select_from(Student)) == 2


def test_class_import_create_and_update(client, importer, db_session, existing_class):
    response = preview(
        client, "classes", [[existing_class.class_id, 12, "12Z"], [None, 10, "10X"]]
    )
    batch, page = batch_for(response, client, db_session)
    assert "Konfirmasi perubahan data Kelas" in page.text
    assert (
        client.post(
            f"/admin/import/{batch.token}/confirm", data={"selected": [2, 3]}
        ).status_code
        == 200
    )
    db_session.refresh(existing_class)
    assert existing_class.class_name == "12Z"
    assert db_session.scalar(select(func.count()).select_from(Class)) == 2


def test_cell_errors_and_valid_rows_can_be_confirmed(client, importer, db_session):
    response = preview(
        client,
        rows=[
            [None, "Valid", None, "0012345678", "Aktif", None],
            [None, "Invalid", None, "123", "TRUE", "bad"],
            [None, "=1+1", None, "0012345679", "Aktif", None],
        ],
    )
    batch, page = batch_for(response, client, db_session)
    assert len(batch.payload["rows"]) == 1
    assert "Sel E3; Sheet students" in page.text
    assert "Sel B4; Sheet students" in page.text
    assert (
        client.post(
            f"/admin/import/{batch.token}/confirm", data={"selected": [2]}
        ).status_code
        == 200
    )
    assert db_session.scalar(select(func.count()).select_from(Student)) == 1


@pytest.mark.parametrize(
    "row,cell",
    [
        ([None, "N", None, "123", "Aktif", None], "D2"),
        ([None, "N", 2147483647, "0012345678", "Aktif", None], "C2"),
        ([2147483647, "N", None, "0012345678", "Aktif", None], "A2"),
        ([None, "N", None, "0012345678", "Aktif", "+62812345"], "F2"),
    ],
)
def test_invalid_rows_report_cells(client, importer, db_session, row, cell):
    batch, page = batch_for(preview(client, rows=[row]), client, db_session)
    assert batch.payload["rows"] == []
    assert f"Sel {cell}" in page.text
    assert "File tidak dapat diproses" in page.text
    assert "data-confirm-import" not in page.text


def test_duplicates_and_empty_rows(client, importer, db_session):
    row = [None, "N", None, "0012345678", "Aktif", None]
    batch, _ = batch_for(
        preview(client, rows=[row, [None] * 6, row]), client, db_session
    )
    assert batch.payload["rows"] == []
    assert batch.payload["total"] == 2
    assert {e["cell"] for e in batch.payload["errors"]} == {"D2", "D4"}


@pytest.mark.parametrize(
    "content,filename,status",
    [
        (b"not a zip", "bad.xlsx", 422),
        (b"data", "bad.csv", 422),
        (b"x" * (import_service.MAX_BYTES + 1), "large.xlsx", 413),
    ],
)
def test_bad_files(client, importer, content, filename, status):
    assert preview(client, content=content, filename=filename).status_code == status


def test_empty_wrong_template_and_multiple_files(client, importer):
    assert preview(client, rows=[]).status_code == 422
    assert (
        preview(
            client, content=workbook_bytes("classes", [[None, 10, "10A"]])
        ).status_code
        == 422
    )
    response = client.post(
        "/admin/import/preview/students",
        files=[("file", ("one.xlsx", b"a")), ("file", ("two.xlsx", b"b"))],
    )
    assert response.status_code == 400


def test_cancel_ownership_and_expiry(client, importer, db_session):
    batch, _ = batch_for(
        preview(client, "classes", [[None, 10, "10X"]]), client, db_session
    )
    with pytest.raises(AppException):
        import_service.get_preview(db_session, importer.id + 1, batch.token)
    assert client.post(f"/admin/import/{batch.token}/cancel").status_code == 200
    assert (
        client.post(
            f"/admin/import/{batch.token}/confirm", data={"selected": [2]}
        ).status_code
        == 409
    )
    batch.expires_at = datetime.now(timezone.utc) - timedelta(seconds=1)
    db_session.commit()
    assert client.get(f"/admin/import?batch={batch.token}").status_code == 404
    assert db_session.scalar(select(func.count()).select_from(Class)) == 0


def test_stale_preview_does_not_overwrite_changes(
    client, importer, db_session, existing_student
):
    batch, _ = batch_for(
        preview(
            client,
            rows=[
                [
                    existing_student.id,
                    "From file",
                    existing_student.class_id,
                    existing_student.nisn,
                    "Aktif",
                    None,
                ]
            ],
        ),
        client,
        db_session,
    )
    existing_student.name = "Newer edit"
    db_session.commit()
    assert (
        client.post(
            f"/admin/import/{batch.token}/confirm", data={"selected": [2]}
        ).status_code
        == 409
    )
    db_session.refresh(existing_student)
    assert existing_student.name == "Newer edit"


def test_atomic_rollback_on_later_write_failure(
    client, importer, db_session, monkeypatch
):
    batch, _ = batch_for(
        preview(
            client,
            rows=[
                [None, "One", None, "0012345678", "Aktif", None],
                [None, "Two", None, "0012345679", "Aktif", None],
            ],
        ),
        client,
        db_session,
    )
    original = student_service.post_student
    calls = 0

    def fail_second(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise AppException("simulated failure", 409)
        return original(*args, **kwargs)

    monkeypatch.setattr(student_service, "post_student", fail_second)
    assert (
        client.post(
            f"/admin/import/{batch.token}/confirm", data={"selected": [2, 3]}
        ).status_code
        == 409
    )
    assert db_session.scalar(select(func.count()).select_from(Student)) == 0
    db_session.refresh(batch)
    assert batch.state == "pending"


def test_tampered_or_empty_selection_cannot_write(client, importer, db_session):
    batch, _ = batch_for(
        preview(client, "classes", [[None, 10, "10X"]]), client, db_session
    )
    for data in ({}, {"selected": [999]}):
        assert (
            client.post(f"/admin/import/{batch.token}/confirm", data=data).status_code
            == 422
        )
    assert db_session.scalar(select(func.count()).select_from(Class)) == 0
