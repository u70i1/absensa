"""Dashboard selection confirmation and atomic bulk actions."""

from io import BytesIO

import pytest
from app.core.admin_auth import ADMIN_SESSION_COOKIE
from app.models.admin import Admin
from app.models.class_ import Class
from app.models.scan_log import ScanLog
from app.models.student import Student
from app.services.admin_auth_service import create_admin_session, hash_password
from app.services.class_service import get_class_directory
from openpyxl import load_workbook


@pytest.fixture
def admin_client(client, db_session):
    admin = Admin(
        username="selection-admin", password_hash=hash_password("test-password")
    )
    db_session.add(admin)
    db_session.commit()
    client.cookies.set(
        ADMIN_SESSION_COOKIE, create_admin_session(db_session, admin, 1), path="/"
    )
    return client


def submit(client, kind, ids, action="delete", step="apply", **kwargs):
    return client.post(
        f"/admin/{kind}/selection/{step}",
        data={"ids": [str(i) for i in ids], "action": action},
        follow_redirects=False,
        **kwargs,
    )


@pytest.mark.parametrize("kind", ["students", "classes"])
@pytest.mark.parametrize("step", ["confirm", "apply"])
def test_requires_admin(client, kind, step):
    response = submit(client, kind, [1], step=step)
    assert response.status_code == 303
    assert response.headers["location"] == "/admin"


def test_student_confirmation_and_subset_delete(
    admin_client, db_session, student_factory
):
    first = student_factory()
    second = student_factory(nisn="1000000002")
    ids = [first.id, second.id]
    preview = submit(
        admin_client,
        "students",
        [first.id],
        step="confirm",
        headers={"HX-Request": "true"},
    )
    assert preview.status_code == 200
    assert "Hapus 1 siswa?" in preview.text and first.nisn in preview.text
    assert 'name="ids"' in preview.text
    assert db_session.get(Student, ids[0]) is not None
    response = submit(
        admin_client, "students", [first.id], headers={"HX-Request": "true"}
    )
    assert response.status_code == 200
    assert response.headers["HX-Retarget"] == "#student-results"
    assert db_session.get(Student, ids[0]) is None
    assert db_session.get(Student, ids[1]) is not None


def test_deactivate_keeps_student_data(admin_client, db_session, student_factory):
    first = student_factory()
    second = student_factory(nisn="1000000002")
    response = submit(
        admin_client, "students", [first.id, first.id], action="deactivate"
    )
    assert response.status_code == 303
    db_session.refresh(first)
    db_session.refresh(second)
    assert first.current is False and second.current is True
    assert first.name == "Shaun" and first.nisn == "1000000001"


def test_bulk_student_delete_keeps_attendance_history(
    admin_client, db_session, student_factory, scan_log_factory
):
    student = student_factory()
    log = scan_log_factory(student)
    student_id, scan_id = student.id, log.scan_id

    assert submit(admin_client, "students", [student_id]).status_code == 303
    db_session.expire_all()
    assert db_session.get(Student, student_id) is None
    assert db_session.get(ScanLog, scan_id).student_id is None


def test_class_delete_unassigns_students(
    admin_client, db_session, class_factory, student_factory
):
    first = class_factory()
    second = class_factory(class_name="Other")
    student = student_factory(class_id=first.class_id)
    first_id = first.class_id
    preview = submit(admin_client, "classes", [first_id], step="confirm")
    assert preview.status_code == 200
    assert "Data siswa tetap tersimpan" in preview.text
    response = submit(admin_client, "classes", [first_id])
    assert response.status_code == 303
    db_session.expire_all()
    assert db_session.get(Class, first_id) is None
    assert db_session.get(Class, second.class_id) is not None
    assert student.class_id is None


@pytest.mark.parametrize("ids", [[], ["abc"], [-1], [999999]])
def test_invalid_selection(admin_client, ids):
    response = submit(admin_client, "students", ids)
    assert response.status_code in {409, 422}


def test_missing_record_prevents_partial_change(
    admin_client, db_session, student_factory
):
    student = student_factory()
    response = submit(
        admin_client, "students", [student.id, 999999], action="deactivate"
    )
    assert response.status_code == 409
    db_session.refresh(student)
    assert student.current


def test_invalid_action_and_csrf(admin_client, class_factory):
    cls = class_factory()
    assert (
        submit(admin_client, "classes", [cls.class_id], action="deactivate").status_code
        == 422
    )
    assert (
        submit(
            admin_client,
            "classes",
            [cls.class_id],
            headers={"Origin": "https://evil.example"},
        ).status_code
        == 403
    )


def test_class_directory_sorted_by_id(db_session, class_factory):
    first = class_factory(class_name="Z", grade=12)
    second = class_factory(class_name="A", grade=1)
    assert [row.class_id for row in get_class_directory(db_session, None, None)] == [
        first.class_id,
        second.class_id,
    ]


def exported_workbook(response):
    assert response.status_code == 200
    assert "attachment;" in response.headers["content-disposition"]
    return load_workbook(BytesIO(response.content))


def exported_student_ids(workbook):
    return {
        row[0].value
        for row in workbook["students"].iter_rows(min_row=2)
        if row[0].value is not None
    }


def test_selected_students_export_keeps_choices_outside_current_filter(
    admin_client, db_session, class_factory, student_factory
):
    first_class = class_factory()
    other_class = class_factory(class_name="Other", grade=11)
    first = student_factory(name="Visible", class_id=first_class.class_id)
    second = student_factory(
        name="Hidden", nisn="1000000002", class_id=other_class.class_id
    )
    omitted = student_factory(name="Omitted", nisn="1000000003")
    data = {"action": "export", "ids": [first.id, second.id, first.id]}
    confirmation = admin_client.post(
        "/admin/students/selection/confirm?q=Visible",
        data=data,
        headers={"HX-Request": "true"},
    )
    assert confirmation.status_code == 200
    assert "Ekspor 2 siswa?" in confirmation.text
    assert 'Pencarian "Visible"' in confirmation.text.replace("&#34;", '"')
    assert second.name in confirmation.text
    assert "data-close-modal>Batal" in confirmation.text
    assert (
        'method="post"' in confirmation.text
        and "data-export-download" in confirmation.text
    )
    assert "content-disposition" not in confirmation.headers

    workbook = exported_workbook(
        admin_client.post("/admin/students/selection/export?q=Visible", data=data)
    )
    assert exported_student_ids(workbook) == {first.id, second.id}
    assert "2 siswa terpilih" in workbook["instructions"]["D6"].value
    assert 'Pencarian "Visible"' in workbook["instructions"]["D6"].value
    directory_ids = {
        row[0].value
        for row in workbook["classes"].iter_rows(min_row=2)
        if isinstance(row[0].value, int)
    }
    assert directory_ids == {first_class.class_id, other_class.class_id}
    assert db_session.get(Student, omitted.id) is not None


def test_selected_classes_export_only_their_students(
    admin_client, class_factory, student_factory
):
    first_class = class_factory(class_name="First", grade=10)
    second_class = class_factory(class_name="Second", grade=11)
    other_class = class_factory(class_name="Other", grade=12)
    first = student_factory(class_id=first_class.class_id)
    second = student_factory(nisn="1000000002", class_id=second_class.class_id)
    student_factory(nisn="1000000003", class_id=other_class.class_id)
    student_factory(nisn="1000000004")
    data = {"action": "export", "ids": [first_class.class_id, second_class.class_id]}
    confirmation = admin_client.post(
        "/admin/classes/selection/confirm?grade=10&class=First",
        data=data,
        headers={"HX-Request": "true"},
    )
    assert confirmation.status_code == 200
    assert "Ekspor 2 siswa?" in confirmation.text
    assert "2 kelas" in confirmation.text and "Jenjang 10" in confirmation.text
    assert "Second" in confirmation.text
    assert "seluruh daftar kelas (3 kelas)" in confirmation.text

    workbook = exported_workbook(
        admin_client.post(
            "/admin/classes/selection/export?grade=10&class=First", data=data
        )
    )
    assert exported_student_ids(workbook) == {first.id, second.id}
    assert "Siswa dari 2 kelas terpilih" in workbook["instructions"]["D6"].value


def test_filtered_export_confirmation_counts_all_matching_pages(
    admin_client, class_factory, student_factory
):
    class_ = class_factory()
    first = student_factory(name="Cut First", class_id=class_.class_id)
    second = student_factory(
        name="Cut Second", nisn="1000000002", class_id=class_.class_id
    )
    student_factory(name="Other", nisn="1000000003", class_id=class_.class_id)
    confirmation = admin_client.get(
        "/admin/students/export/confirm?q=Cut&grade=10&page=2",
        headers={"HX-Request": "true"},
    )
    assert confirmation.status_code == 200
    assert confirmation.headers["HX-Retarget"] == "#modal-content"
    assert "Ekspor 2 siswa?" in confirmation.text
    assert "Jenjang 10" in confirmation.text
    assert 'name="q" value="Cut"' in confirmation.text
    assert 'name="grade" value="10"' in confirmation.text
    assert 'name="page"' not in confirmation.text
    assert "data-close-modal>Batal" in confirmation.text
    workbook = exported_workbook(
        admin_client.get("/admin/students/export?q=Cut&grade=10&page=2")
    )
    assert exported_student_ids(workbook) == {first.id, second.id}
    dashboard = admin_client.get("/admin/students?q=Cut")
    assert "/admin/students/export/confirm?" in dashboard.text
    assert "Ekspor terpilih" in dashboard.text


@pytest.mark.parametrize("kind", ["students", "classes"])
@pytest.mark.parametrize("ids", [[], ["bad"], [-1], [999999]])
def test_selected_export_rejects_invalid_ids(admin_client, kind, ids):
    response = submit(admin_client, kind, ids, action="export", step="export")
    assert response.status_code in {409, 422}
    assert "content-disposition" not in response.headers


def test_export_cannot_be_applied_as_a_mutation(
    admin_client, db_session, student_factory
):
    student = student_factory()
    response = submit(
        admin_client, "students", [student.id], action="export", step="apply"
    )
    assert response.status_code == 422
    assert db_session.get(Student, student.id) is not None


def test_empty_selected_class_export(admin_client, class_factory):
    class_ = class_factory()
    confirmation = submit(
        admin_client, "classes", [class_.class_id], action="export", step="confirm"
    )
    assert "Ekspor 0 siswa?" in confirmation.text
    workbook = exported_workbook(
        submit(
            admin_client, "classes", [class_.class_id], action="export", step="export"
        )
    )
    assert exported_student_ids(workbook) == set()
    assert workbook["classes"]["A2"].value == class_.class_id


def test_export_confirmation_and_download_require_admin(client):
    assert (
        client.get("/admin/students/export/confirm", follow_redirects=False).status_code
        == 303
    )
    for kind in ("students", "classes"):
        response = submit(client, kind, [1], action="export", step="export")
        assert response.status_code == 303
