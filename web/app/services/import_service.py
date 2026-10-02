"""Review spreadsheets and archive photos, then apply entire workbooks atomically."""

import re
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from secrets import token_urlsafe
from typing import Any, BinaryIO
from zipfile import BadZipFile, ZipFile

from app.models.class_ import Class
from app.models.import_batch import ImportBatch, ImportPhoto
from app.models.student import Student
from app.schemas.import_ import ClassImportRow, StudentImportRow
from app.services import (
    import_archive_service,
    student_photo_service,
)
from app.services.exceptions import AppException
from app.services.export_service import (
    CLASS_COLUMNS,
    STUDENT_COLUMNS,
    class_lookup_formula,
)
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import ValidationError
from sqlalchemy import delete, select, text
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import Session

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10000
COLUMNS = {"students": STUDENT_COLUMNS, "classes": CLASS_COLUMNS}
FIELD_ERRORS = {
    "id": "ID harus berupa bilangan bulat positif, atau kosong untuk data baru.",
    "class_id": "ID KELAS harus berupa ID angka atau ID sementara seperti N1 yang tercantum pada sheet classes.",
    "name": "NAMA wajib diisi, maksimal 255 karakter.",
    "nisn": "NISN wajib berisi tepat 10 digit. Gunakan format teks untuk mempertahankan nol di depan.",
    "current": "STATUS wajib diisi dengan Aktif atau Tidak aktif.",
    "guardian_phone": "NOMOR WALI harus berisi angka saja, diawali 08 atau 628, maksimal 32 digit.",
    "grade": "JENJANG wajib berupa bilangan bulat antara 1 dan 20.",
    "class_name": "NAMA KELAS wajib diisi, maksimal 20 karakter.",
}


def _error(kind, row, field, message, sheet_name=None):
    column = next(i for i, (_, key) in enumerate(COLUMNS[kind], 1) if key == field)
    return {
        "cell": f"{row}:{get_column_letter(column)}",
        "sheet": sheet_name or kind,
        "message": message,
    }


def _value(field, value):
    if isinstance(value, str):
        value = value.strip()
    if value in (None, ""):
        return None
    if field == "current":
        if value not in ("Aktif", "Tidak aktif"):
            raise ValueError(FIELD_ERRORS[field])
        return value == "Aktif"
    if (
        field == "class_id"
        and isinstance(value, str)
        and re.fullmatch(r"N[1-9][0-9]*", value)
    ):
        return value
    if field in {"id", "class_id", "grade"}:
        if isinstance(value, bool):
            raise ValueError(FIELD_ERRORS[field])
        if isinstance(value, float) and value.is_integer():
            return int(value)
        if isinstance(value, int) or (
            isinstance(value, str) and value.isascii() and value.isdigit()
        ):
            return int(value)
        raise ValueError(FIELD_ERRORS[field])
    if not isinstance(value, str):
        # Never guess zeros already lost by Excel's numeric conversion.
        if (
            field in {"nisn", "guardian_phone"}
            and isinstance(value, (int, float))
            and not isinstance(value, bool)
            and int(value) == value
        ):
            return str(int(value))
        raise ValueError(FIELD_ERRORS[field])
    return value


def _is_class_lookup(value, row: int, result_column: int) -> bool:
    if not isinstance(value, str):
        return False
    # Calc serializes FALSE as FALSE(); both are standard XLSX formulas.
    normalized = (
        re.sub(r"\s+", "", value).upper().replace("FALSE()", "FALSE").replace("$", "")
    )
    return normalized == class_lookup_formula(row, result_column).upper()


def _read_rows(content: bytes):
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 2000
                or sum(item.file_size for item in entries) > 50 * 1024 * 1024
            ):
                raise AppException(
                    "Isi file terlalu besar setelah dibuka. Pecah data menjadi file yang lebih kecil.",
                    413,
                )
            if any("vbaproject" in item.filename.lower() for item in entries):
                raise AppException("Gunakan file .xlsx tanpa makro.", 422)
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except AppException:
        raise
    except Exception as exc:
        raise AppException(
            "File Excel rusak atau bukan file .xlsx yang valid.", 422
        ) from exc
    try:
        parsed, errors, used_rows = [], [], 0
        for required in ("instructions", "students", "classes"):
            if required not in workbook.sheetnames:
                errors.append(
                    {
                        "sheet": required,
                        "cell": "1:A",
                        "message": "Sheet wajib tidak ditemukan. Jangan mengganti nama sheet.",
                    }
                )
        if not any(name in workbook.sheetnames for name in COLUMNS):
            raise AppException(
                " | ".join(
                    f"{e['sheet']} > {e['cell']} > {e['message']}" for e in errors
                ),
                422,
            )
        sheet_specs = []
        if "classes" in workbook.sheetnames:
            sheet_specs.append(("classes", "classes", 0))
        student_sheet_names = []
        if "students" in workbook.sheetnames:
            student_sheet_names.append("students")
        # A copied student worksheet may use any tab name. Recognize it from
        # its first two headers, then validate the complete header normally.
        for name in workbook.sheetnames:
            if name in {"instructions", "students", "classes"}:
                continue
            candidate = workbook[name]
            first_cells = next(candidate.iter_rows(max_row=1, max_col=2), ())
            if len(first_cells) == 2 and tuple(c.value for c in first_cells) == (
                STUDENT_COLUMNS[0][0],
                STUDENT_COLUMNS[1][0],
            ):
                student_sheet_names.append(name)
        sheet_specs.extend(
            ("students", name, index * 100001)
            for index, name in enumerate(student_sheet_names, 1)
        )
        for kind, sheet_name, sheet_offset in sheet_specs:
            sheet, columns = workbook[sheet_name], COLUMNS[kind]
            if (sheet.max_row or 0) > 100000 or (sheet.max_column or 0) > 100:
                raise AppException(
                    f"{sheet_name} > 1:A > Ukuran sheet melebihi batas template.", 422
                )
            rows = sheet.iter_rows(max_col=len(columns))
            header = next(rows, ())
            header_errors = [
                _error(
                    kind,
                    1,
                    field,
                    f'Kolom wajib "{name}" tidak ditemukan atau diganti. Gunakan urutan kolom template.',
                    sheet_name,
                )
                for (name, field), cell in zip(columns, header)
                if cell.value != name
            ]
            errors.extend(header_errors)
            if header_errors:
                continue
            schema = StudentImportRow if kind == "students" else ClassImportRow
            seen_class_ids = set()
            for number, cells in enumerate(rows, 2):
                if kind == "classes" and cells[0].value not in (None, ""):
                    reference = str(cells[0].value).strip()
                    if reference in seen_class_ids:
                        errors.append(
                            _error(
                                kind,
                                number,
                                "class_id",
                                "ID kelas berulang dalam sheet, termasuk baris cadangan. Setiap ID harus unik.",
                                sheet_name,
                            )
                        )
                    seen_class_ids.add(reference)
                if (
                    kind == "classes"
                    and isinstance(cells[0].value, str)
                    and re.fullmatch(r"N[1-9][0-9]*", cells[0].value)
                    and all(c.value in (None, "") for c in cells[1:])
                ):
                    continue
                if all(
                    cell.value is None
                    or (isinstance(cell.value, str) and not cell.value.strip())
                    or (
                        kind == "students"
                        and index in (6, 7)
                        and _is_class_lookup(cell.value, number, index - 4)
                    )
                    for index, cell in enumerate(cells)
                ):
                    continue
                used_rows += 1
                if used_rows > MAX_ROWS:
                    raise AppException("Maksimal 10.000 baris data per file.", 413)
                values: dict[str, Any] = {}
                row_errors = []
                for (_, field), cell in zip(columns, cells):
                    try:
                        if kind == "students" and field in {"grade", "class_name"}:
                            if cell.value not in (None, "") and not _is_class_lookup(
                                cell.value, number, 2 if field == "grade" else 3
                            ):
                                raise ValueError(
                                    "Kolom ini hanya boleh berisi rumus bawaan. Ubah jenjang/nama kelas pada sheet classes, bukan students."
                                )
                            values[field] = None
                            continue
                        if cell.data_type in {"f", "e"}:
                            raise ValueError(
                                "Gunakan nilai langsung, bukan rumus atau nilai kesalahan Excel."
                            )
                        values[field] = _value(field, cell.value)
                    except ValueError as exc:
                        row_errors.append(
                            _error(kind, number, field, str(exc), sheet_name)
                        )
                # Validate independent fields even when another cell failed parsing.
                try:
                    validated = schema.model_validate(values).model_dump()
                    if kind == "students":
                        validated.setdefault("guardian_phone", None)
                    values = validated
                except ValidationError as exc:
                    row_errors.extend(
                        _error(
                            kind,
                            number,
                            str(e["loc"][0]),
                            FIELD_ERRORS[str(e["loc"][0])],
                            sheet_name,
                        )
                        for e in exc.errors()
                    )
                errors.extend(row_errors)
                if not row_errors:
                    parsed.append(
                        {
                            "row": number,
                            "kind": kind,
                            "sheet": sheet_name,
                            "sheet_offset": sheet_offset,
                            "values": values,
                        }
                    )
        if not used_rows and not errors:
            raise AppException(
                "File tidak berisi data. Isi baris di bawah judul kolom terlebih dahulu.",
                422,
            )
        return parsed, errors, used_rows, max(1, len(student_sheet_names) + 1)
    except AppException:
        raise
    except Exception as exc:
        raise AppException(
            "Isi sheet tidak dapat dibaca. Simpan ulang sebagai .xlsx.", 422
        ) from exc
    finally:
        workbook.close()


STUDENT_FIELDS = ("name", "class_id", "nisn", "current", "guardian_phone")


def _reference_key(row, value):
    # Temporary references are local to each workbook in an archive.
    return (row.get("file", ""), value) if isinstance(value, str) else value


def _changed_fields(row, before):
    fields = STUDENT_FIELDS if row["kind"] == "students" else ("grade", "class_name")
    changed = []
    for field in fields:
        value = (
            row["resolved_class_id"]
            if row["kind"] == "students" and field == "class_id"
            else row["values"][field]
        )
        if (
            before is None
            or value != before[field]
            or (
                row["kind"] == "students"
                and field == "class_id"
                and row["class_key"] is not None
                and value is None
            )
        ):
            changed.append(field)
    if row.get("incoming_photo"):
        changed.append("photo_path")
    return changed


def _action(before, changed):
    return "create" if before is None else "update" if changed else "unchanged"


def _review_workbook(db: Session, rows: list[dict]):
    """Plan against the final class and student state without mutating ORM objects."""
    classes = {item.class_id: item for item in db.scalars(select(Class))}
    students = {item.id: item for item in db.scalars(select(Student))}
    final_classes = {
        key: (item.grade, item.class_name) for key, item in classes.items()
    }
    class_rows = [r for r in rows if r["kind"] == "classes"]
    student_rows = [r for r in rows if r["kind"] == "students"]
    errors, invalid = [], set()

    def error(row, field, message):
        errors.append(
            {
                **_error(row["kind"], row["row"], field, message, row.get("sheet")),
                "file": row.get("file", ""),
            }
        )
        invalid.add(id(row))

    for group, existing, id_field in (
        (class_rows, classes, "class_id"),
        (student_rows, students, "id"),
    ):
        references = defaultdict(list)
        for row in group:
            references[_reference_key(row, row["values"][id_field])].append(row)
        for row in group:
            value = row["values"][id_field]
            if isinstance(value, int) and value not in existing:
                error(
                    row,
                    id_field,
                    f"ID {value} tidak ditemukan. Jangan mengubah ID yang sudah ada; kosongkan ID untuk data baru.",
                )
            declarations = references[_reference_key(row, value)]
            shared_class = (
                id_field == "class_id"
                and isinstance(value, int)
                and len({r.get("file", "") for r in declarations}) == len(declarations)
                and len(
                    {
                        (r["values"]["grade"], r["values"]["class_name"])
                        for r in declarations
                    }
                )
                == 1
            )
            if value is not None and len(declarations) > 1 and not shared_class:
                error(
                    row,
                    id_field,
                    "ID muncul lebih dari sekali dalam workbook/unggahan.",
                )

    # Existing IDs describe edits. New references describe a class pair and may
    # share it across workbooks or reuse a class already in the final state.
    for row in class_rows:
        values = row["values"]
        if isinstance(values["class_id"], int):
            final_classes[values["class_id"]] = (
                values["grade"],
                values["class_name"],
            )
    final_class_ids = {pair: key for key, pair in final_classes.items()}
    declared_classes = {
        (row.get("file", ""), row["values"]["class_id"])
        if row["values"]["class_id"] is not None
        else f"new:{index}": (row["values"]["grade"], row["values"]["class_name"])
        for index, row in enumerate(class_rows)
    }
    pair_counts = Counter(final_classes.values())
    for row in class_rows:
        values = row["values"]
        pair = (values["grade"], values["class_name"])
        if pair_counts[pair] > 1:
            error(
                row,
                "class_name",
                f'Kelas "{pair[1]}" pada jenjang {pair[0]} sudah terdaftar atau berulang dalam workbook.',
            )

    final_nisns: dict[int | str, str] = {
        key: item.nisn for key, item in students.items()
    }
    for index, row in enumerate(student_rows):
        values = row["values"]
        final_nisns[values["id"] if values["id"] is not None else f"new:{index}"] = (
            values["nisn"]
        )
    nisn_counts = Counter(final_nisns.values())
    for row in student_rows:
        values = row["values"]
        if nisn_counts[values["nisn"]] > 1:
            error(row, "nisn", "NISN sudah digunakan atau berulang dalam workbook.")
        reference = values["class_id"]
        student_class_pair = (
            declared_classes.get((row.get("file", ""), reference))
            if reference is not None
            else None
        )
        if reference is not None and (
            student_class_pair is None
            or (isinstance(reference, int) and reference not in classes)
        ):
            error(
                row,
                "class_id",
                f"ID kelas {reference} tidak ditemukan atau belum lengkap pada sheet classes. Isi kelas tersebut terlebih dahulu.",
            )
        if reference is None and values["id"] is None:
            error(
                row,
                "class_id",
                "ID KELAS wajib diisi untuk siswa baru. Gunakan ID dari sheet classes.",
            )
        row["class_key"] = (
            list(student_class_pair) if student_class_pair is not None else None
        )
        row["resolved_class_id"] = (
            reference
            if isinstance(reference, int)
            else final_class_ids.get(student_class_pair)
            if student_class_pair is not None
            else None
        )
        row["class_name"] = (
            student_class_pair[1] if student_class_pair is not None else "—"
        )

    valid, seen_classes = [], set()
    for row in rows:
        if id(row) in invalid:
            continue
        kind, values = row["kind"], row["values"]
        id_field, existing, fields = (
            ("id", students, STUDENT_FIELDS)
            if kind == "students"
            else ("class_id", classes, ("grade", "class_name"))
        )
        item = existing.get(values[id_field])
        before = {field: getattr(item, field) for field in fields} if item else None
        if (
            before is not None
            and isinstance(item, Student)
            and row.get("incoming_photo")
        ):
            before["photo_path"] = item.photo_path
        changed = _changed_fields(row, before)
        reviewed = {
            **row,
            "before": before,
            "changed": changed,
            "action": _action(before, changed),
        }
        if kind == "classes":
            pair = (values["grade"], values["class_name"])
            class_key = values["class_id"] if item else pair
            if class_key in seen_classes or (item is None and pair in final_class_ids):
                reviewed["action"] = "reuse"
                reviewed["changed"] = []
            seen_classes.add(class_key)
        if kind == "students":
            reviewed["photo_path"] = (
                item.photo_path if isinstance(item, Student) else None
            )
        valid.append(reviewed)
    return valid, errors


def create_preview(
    db: Session, admin_id: int, kind: str | None, filename: str, source: BinaryIO
) -> ImportBatch:
    if kind not in (None, "students"):
        raise AppException("Jenis impor tidak ditemukan.", 404)
    archive = import_archive_service.is_archive(filename)
    if not archive and Path(filename).suffix.lower() != ".xlsx":
        raise AppException(
            "Pilih file .xlsx, .zip, .rar, .7z, .tar, atau TAR terkompresi.", 422
        )
    limit = import_archive_service.MAX_ARCHIVE_BYTES if archive else MAX_BYTES
    content = source.read(limit + 1)
    if len(content) > limit:
        raise AppException(
            "Ukuran file maksimal 100 MB untuk arsip atau 10 MB untuk Excel.", 413
        )
    filename = filename.replace("\\", "/").rsplit("/", 1)[-1][:255]
    workbooks, photos = (
        import_archive_service.read_archive(content)
        if archive
        else ([(filename, content)], {})
    )
    all_rows = []
    errors: list[dict] = []
    files = []
    total = offset = expanded = 0
    for name, workbook_content in workbooks:
        try:
            # Bound the combined XML expansion of all XLSX members as well.
            try:
                with ZipFile(BytesIO(workbook_content)) as workbook_zip:
                    expanded += sum(
                        entry.file_size for entry in workbook_zip.infolist()
                    )
            except BadZipFile as exc:
                raise AppException(
                    "File Excel rusak atau bukan file .xlsx yang valid.", 422
                ) from exc
            if expanded > import_archive_service.MAX_EXPANDED_BYTES:
                raise AppException(
                    "Total isi workbook terlalu besar setelah dibuka (maksimal 200 MB).",
                    413,
                )
            parsed, file_errors, used, sheet_slots = _read_rows(workbook_content)
            total += used
            if total > MAX_ROWS:
                raise AppException(
                    "Maksimal 10.000 baris data untuk seluruh unggahan.", 413
                )
            for row in parsed:
                row.update(
                    file=name,
                    key=offset + row["sheet_offset"] + row["row"],
                )
            all_rows.extend(parsed)
            errors.extend({**error, "file": name} for error in file_errors)
            files.append({"name": name, "kind": "mixed", "rows": used})
            # Each sheet is bounded at 100,000 rows, so keys cannot collide.
            offset += sheet_slots * 100001
        except AppException as exc:
            if not archive or exc.status_code == 413:
                raise
            errors.append(
                {"file": name, "cell": None, "sheet": None, "message": exc.detail}
            )
    valid, review_errors = _review_workbook(db, all_rows)
    errors.extend(review_errors)
    staged_photos = {}
    students = {
        row["values"]["nisn"]: row for row in valid if row["kind"] == "students"
    }
    photo_bytes = 0
    for nisn, (name, image_content) in photos.items():
        try:
            if nisn not in students:
                raise AppException(
                    "Foto tidak memiliki baris siswa yang valid dengan NISN tersebut dalam unggahan ini.",
                    422,
                )
            normalized = student_photo_service.normalize_photo(BytesIO(image_content))
            photo_bytes += len(normalized)
            if photo_bytes > 50 * 1024 * 1024:
                raise AppException("Total foto setelah diproses maksimal 50 MB.", 413)
            existing_path = students[nisn]["photo_path"]
            if existing_path:
                try:
                    with student_photo_service.photo_file(existing_path).open(
                        "rb"
                    ) as saved:
                        if saved.read(len(normalized) + 1) == normalized:
                            continue
                except (OSError, AppException):
                    pass  # Missing photos should be restored from the upload.
            staged_photos[nisn] = normalized
            students[nisn]["incoming_photo"] = name
            students[nisn]["changed"].append("photo_path")
            if students[nisn]["before"] is not None:
                students[nisn]["before"]["photo_path"] = students[nisn]["photo_path"]
                students[nisn]["action"] = "update"
        except AppException as exc:
            if exc.status_code == 413 and photo_bytes > 50 * 1024 * 1024:
                raise
            errors.append(
                {"file": name, "cell": None, "sheet": "photos", "message": exc.detail}
            )
    now = datetime.now(timezone.utc)
    db.execute(delete(ImportBatch).where(ImportBatch.expires_at < now))
    detected_kinds = {item["kind"] for item in files}
    batch_kind = next(iter(detected_kinds)) if len(detected_kinds) == 1 else "mixed"
    batch = ImportBatch(
        token=token_urlsafe(32),
        admin_id=admin_id,
        kind=batch_kind,
        filename=filename,
        size=len(content),
        payload={
            "rows": valid,
            "errors": errors,
            "total": total,
            "files": files,
            "photo_count": len(staged_photos),
        },
        state="pending",
        expires_at=now + timedelta(hours=1),
    )
    db.add(batch)
    db.flush()
    db.add_all(
        ImportPhoto(batch_token=batch.token, nisn=nisn, content=data)
        for nisn, data in staged_photos.items()
    )
    db.commit()
    return batch


def get_preview(db: Session, admin_id: int, token: str, *, lock=False) -> ImportBatch:
    stmt = select(ImportBatch).where(
        ImportBatch.token == token, ImportBatch.admin_id == admin_id
    )
    if lock:
        stmt = stmt.with_for_update().execution_options(populate_existing=True)
    batch = db.scalar(stmt)
    if batch is None or batch.expires_at <= datetime.now(timezone.utc):
        raise AppException(
            "Pratinjau tidak ditemukan atau sudah kedaluwarsa. Unggah file kembali.",
            404,
        )
    return batch


def cancel_preview(db: Session, admin_id: int, token: str):
    batch = get_preview(db, admin_id, token, lock=True)
    if batch.state == "pending":
        batch.state = "cancelled"
        batch.payload = {}
        db.execute(delete(ImportPhoto).where(ImportPhoto.batch_token == token))
        db.commit()


def get_preview_photo(db: Session, admin_id: int, token: str, nisn: str) -> bytes:
    batch = get_preview(db, admin_id, token)
    photo = db.get(ImportPhoto, (token, nisn)) if batch.state == "pending" else None
    if photo is None:
        raise AppException("Foto pratinjau tidak ditemukan.", 404)
    return photo.content


def get_editable_row(db: Session, admin_id: int, token: str, key: int, *, lock=False):
    batch = get_preview(db, admin_id, token, lock=lock)
    if batch.state != "pending":
        raise AppException(
            "Hanya pratinjau yang belum dikonfirmasi yang dapat diedit.", 409
        )
    row = next(
        (row for row in batch.payload["rows"] if row.get("key", row["row"]) == key),
        None,
    )
    if row is None:
        raise AppException("Baris pratinjau tidak ditemukan.", 404)
    return batch, row


def edit_preview_row(
    db: Session, admin_id: int, token: str, key: int, values: dict, revision: int
):
    """Save a validated draft, retaining the original identity and stale-data snapshot."""
    batch, original = get_editable_row(db, admin_id, token, key, lock=True)
    if revision != original.get("revision", 0):
        raise AppException(
            "Baris ini sudah diedit. Tutup lalu buka kembali editor untuk melihat perubahan terbaru.",
            409,
        )
    kind = original.get("kind", batch.kind)
    id_field = "id" if kind == "students" else "class_id"
    fields = [
        field
        for _, field in COLUMNS[kind]
        if field != id_field
        and not (kind == "students" and field in {"grade", "class_name"})
    ]
    try:
        parsed = {field: _value(field, values.get(field)) for field in fields}
        parsed[id_field] = original["values"][id_field]
        if kind == "students":
            parsed["grade"] = None
            parsed["class_name"] = None
        schema = StudentImportRow if kind == "students" else ClassImportRow
        parsed = schema.model_validate(parsed).model_dump()
        if kind == "students":
            parsed.setdefault("guardian_phone", None)
    except ValidationError as exc:
        raise AppException(
            " ".join(FIELD_ERRORS[str(error["loc"][0])] for error in exc.errors()), 422
        ) from exc
    except ValueError as exc:
        raise AppException(str(exc), 422) from exc
    payload = deepcopy(batch.payload)
    candidate = next(
        row for row in payload["rows"] if row.get("key", row["row"]) == key
    )
    candidate["values"] = parsed
    reviewed, errors = _review_workbook(db, payload["rows"])
    if errors:
        raise AppException(errors[0]["message"], 422)
    # Do not refresh 'before': editing a draft must never bypass stale-write checks.
    for draft, refreshed in zip(payload["rows"], reviewed):
        if draft["kind"] == "students":
            draft["class_name"] = refreshed["class_name"]
            draft["class_key"] = refreshed["class_key"]
            draft["resolved_class_id"] = refreshed["resolved_class_id"]
        draft["changed"] = _changed_fields(draft, draft["before"])
        draft["action"] = (
            "reuse"
            if refreshed["action"] == "reuse"
            else _action(draft["before"], draft["changed"])
        )
        if draft["action"] == "reuse":
            draft["changed"] = []
    if original.get("incoming_photo"):
        candidate["photo_nisn"] = original.get("photo_nisn", original["values"]["nisn"])
    candidate["revision"] = revision + 1
    batch.payload = payload
    db.commit()
    return batch


def apply_preview(
    db: Session, admin_id: int, token: str, selected: list[int]
) -> ImportBatch:
    batch = get_preview(db, admin_id, token, lock=True)
    if batch.state == "applied":
        return batch  # A retry never creates duplicate records.
    if batch.state != "pending":
        raise AppException("Impor ini sudah dibatalkan. Unggah file kembali.", 409)
    if batch.payload["errors"]:
        raise AppException(
            "Impor ditolak: perbaiki seluruh kesalahan lalu unggah kembali. Tidak ada siswa atau kelas yang disimpan.",
            422,
        )
    rows = batch.payload["rows"]
    if not rows or set(selected) != {row["key"] for row in rows}:
        raise AppException(
            "Impor harus mencakup seluruh baris workbook. Tidak ada perubahan disimpan.",
            422,
        )
    chosen = deepcopy(rows)
    groups = {
        kind: [row for row in chosen if row.get("kind", batch.kind) == kind]
        for kind in ("classes", "students")
    }
    new_photos: list[str] = []
    old_photos: list[str | None] = []
    try:
        # Lock existing rows in a stable order and reject stale previews.
        for kind, group in groups.items():
            id_field, model = (
                ("id", Student) if kind == "students" else ("class_id", Class)
            )
            ids = sorted(
                row["values"][id_field] for row in group if row["before"] is not None
            )
            if kind == "classes":
                ids = sorted(
                    set(ids)
                    | {
                        r["values"]["class_id"]
                        for r in groups["students"]
                        if isinstance(r["values"]["class_id"], int)
                    }
                )
            locked = {
                getattr(item, id_field): item
                for item in db.scalars(
                    select(model)
                    .where(getattr(model, id_field).in_(ids))
                    .order_by(getattr(model, id_field))
                    .with_for_update()
                    .execution_options(populate_existing=True)
                )
            }
            for row in group:
                if row["before"] is not None:
                    item = locked.get(row["values"][id_field])
                    if item is None or any(
                        getattr(item, field) != value
                        for field, value in row["before"].items()
                    ):
                        raise AppException(
                            "Data berubah sejak pratinjau dibuat. Batalkan lalu unggah kembali untuk meninjau perubahan terbaru.",
                            409,
                        )
        chosen, errors = _review_workbook(db, chosen)
        if errors:
            raise AppException(
                "Data tidak lagi valid: "
                + " | ".join(
                    f"{e['sheet']} > {e['cell']} > {e['message']}" for e in errors
                )
                + " Unggah file kembali.",
                409,
            )
        groups = {
            kind: [row for row in chosen if row["kind"] == kind]
            for kind in ("classes", "students")
        }
        # Both unique constraints describe the final state, including swaps.
        db.execute(
            text("SET CONSTRAINTS uq_grade_class_name, students_nisn_key DEFERRED")
        )
        for row in groups["classes"]:
            if row["action"] not in {"create", "update"}:
                continue
            values = row["values"]
            item = (
                db.get(Class, values["class_id"])
                if isinstance(values["class_id"], int)
                else Class()
            )
            if item is None:
                raise AppException(
                    "Kelas berubah sejak pratinjau dibuat. Unggah file kembali.", 409
                )
            item.grade, item.class_name = values["grade"], values["class_name"]
            db.add(item)
        db.flush()
        class_ids = {
            (item.grade, item.class_name): item.class_id
            for item in db.scalars(select(Class))
        }
        for row in groups["students"]:
            if row["action"] not in {"create", "update"}:
                continue
            pair = tuple(row["class_key"]) if row["class_key"] is not None else None
            values = {field: row["values"][field] for field in STUDENT_FIELDS}
            values["class_id"] = class_ids[pair] if pair is not None else None
            student = _apply_student(db, row["values"]["id"], values)
            if row.get("incoming_photo"):
                photo = db.get(
                    ImportPhoto, (token, row.get("photo_nisn", values["nisn"]))
                )
                if photo is None:
                    raise AppException(
                        "Foto pratinjau tidak tersedia. Unggah arsip kembali.", 409
                    )
                path = student_photo_service.store_normalized_photo(photo.content)
                new_photos.append(path)
                old_photos.append(student.photo_path)
                student.photo_path = path
        db.flush()
        db.execute(
            text("SET CONSTRAINTS uq_grade_class_name, students_nisn_key IMMEDIATE")
        )
        batch.state = "applied"
        batch.payload = {
            "created": sum(row["action"] == "create" for row in chosen),
            "updated": sum(row["action"] == "update" for row in chosen),
            "skipped": batch.payload["total"]
            - sum(row["action"] in {"create", "update"} for row in chosen),
            "photos_saved": len(new_photos),
        }
        db.execute(delete(ImportPhoto).where(ImportPhoto.batch_token == token))
        db.commit()
    except Exception as exc:
        db.rollback()
        for path in new_photos:
            student_photo_service.remove_photo(path)
        if isinstance(exc, IntegrityError):
            raise AppException(
                "Data bertabrakan dengan perubahan lain. Tidak ada perubahan disimpan. Unggah file kembali.",
                409,
            ) from exc
        if isinstance(exc, OSError):
            raise AppException(
                "Foto gagal disimpan. Tidak ada perubahan diterapkan. Silakan coba lagi.",
                503,
            ) from exc
        if isinstance(exc, SQLAlchemyError):
            raise AppException(
                "Penyimpanan gagal. Tidak ada perubahan siswa atau kelas yang disimpan. Silakan coba lagi.",
                503,
            ) from exc
        raise
    for old_path in old_photos:
        student_photo_service.remove_photo(old_path)
    return batch


def _apply_student(db: Session, student_id: int | None, values: dict) -> Student:
    """Apply already validated values without per-row commits or uniqueness checks."""
    student = db.get(Student, student_id) if student_id else Student()
    if student is None:
        raise AppException(
            "Siswa berubah sejak pratinjau dibuat. Unggah file kembali.", 409
        )
    for field, value in values.items():
        setattr(student, field, value)
    db.add(student)
    db.flush()
    return student
